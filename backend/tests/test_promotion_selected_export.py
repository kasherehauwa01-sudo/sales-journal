import asyncio
from io import BytesIO
from zipfile import ZipFile
from pathlib import Path
import time
import pytest
from PIL import Image
from pydantic import ValidationError
from openpyxl import load_workbook
from fastapi import HTTPException
from app.services import promotion_selected_export as export
from app.routers import promotion_candidates_report as router


def row(**overrides):
    return {
        "key": "code:001",
        "name": "Товар",
        "code": "001",
        "image_url": "https://catalog.test/a.png",
        "stock": 10,
        "aviatorov_stock": 2,
        **overrides,
    }


def test_selected_workbook_all_columns_and_embedded_photo(monkeypatch):
    calls = []

    def photo(value, directory, index, deadline):
        calls.append(value)
        path = Path(directory) / f"{index}.png"
        Image.new("RGB", (40, 30), "red").save(path)
        return path

    monkeypatch.setattr(export, "thumbnail", photo)
    monkeypatch.setattr(export, "ensure_headroom", lambda: None)
    result = export.excel_selected(export.SelectedExport(rows=[row(name="=1+1")]))
    with ZipFile(BytesIO(result.body)) as archive:
        assert (
            len([name for name in archive.namelist() if name.startswith("xl/media/")])
            == 1
        )
    book = load_workbook(BytesIO(result.body))
    try:
        sheet = book["Выбранные товары"]
        assert [cell.value for cell in sheet[1]] == export.HEADERS
        assert sheet.max_row == 2
        assert sheet["B2"].value == "=1+1" and sheet["B2"].data_type == "s"
        assert sheet["C2"].value == "001" and sheet["M2"].value == 2
        assert len(sheet._images) == 1
    finally:
        book.close()
    assert calls == ["https://catalog.test/a.png"]


def test_missing_photo_preserves_public_link(monkeypatch):
    monkeypatch.setattr(export, "thumbnail", lambda *args: None)
    monkeypatch.setattr(export, "ensure_headroom", lambda: None)
    result = export.excel_selected(
        export.SelectedExport(rows=[row(), row(key="code:002", image_url=None)])
    )
    book = load_workbook(BytesIO(result.body), read_only=True)
    try:
        assert book["Выбранные товары"]["A2"].value == "https://catalog.test/a.png"
        assert book["Выбранные товары"]["A3"].value == "Нет фото"
    finally:
        book.close()


def test_untrusted_origin_and_expired_budget_never_download(monkeypatch, tmp_path):
    monkeypatch.setattr(
        export.settings, "vrcatalog_api_url", "https://catalog.test/api"
    )

    def reject(*args):
        raise AssertionError("network must not run")

    monkeypatch.setattr(export, "build_opener", reject)
    assert (
        export.thumbnail(
            "http://127.0.0.1/secret", tmp_path, 1, time.perf_counter() + 10
        )
        is None
    )
    assert (
        export.thumbnail(
            "https://catalog.test/a.jpg", tmp_path, 1, time.perf_counter() - 1
        )
        is None
    )
    assert export.trusted_image_url("https://user:secret@catalog.test/a.jpg") is None
    assert export.trusted_image_url("https://other.test/a.jpg") is None


def test_selection_limits_and_memory_rejection(monkeypatch):
    with pytest.raises(ValidationError):
        export.SelectedExport(rows=[])
    with pytest.raises(ValidationError):
        export.SelectedExport(rows=[row()] * 301)

    def reject():
        raise HTTPException(503, "memory_headroom")

    monkeypatch.setattr(export, "ensure_headroom", reject)
    with pytest.raises(HTTPException):
        export.excel_selected(export.SelectedExport(rows=[row()]))


def test_thumbnail_download_is_bounded_and_resized(monkeypatch, tmp_path):
    monkeypatch.setattr(
        export.settings, "vrcatalog_api_url", "https://catalog.test/api"
    )
    monkeypatch.setattr(export, "ensure_headroom", lambda: None)
    data = BytesIO()
    Image.new("RGB", (200, 100), "blue").save(data, format="PNG")

    class Opener:
        def open(self, *args, **kwargs):
            return BytesIO(data.getvalue())

    monkeypatch.setattr(export, "build_opener", lambda *args: Opener())
    path = export.thumbnail(
        "https://catalog.test/a.png", tmp_path, 2, time.perf_counter() + 10
    )
    with Image.open(path) as image:
        assert image.size == (96, 48)
    monkeypatch.setattr(export, "MAX_IMAGE_BYTES", 1)
    assert (
        export.thumbnail(
            "https://catalog.test/a.png", tmp_path, 3, time.perf_counter() + 10
        )
        is None
    )
    assert not (tmp_path / "source-image").exists()


def test_endpoint_exports_snapshot_without_dataset_query(monkeypatch):
    async def reject(*args):
        raise AssertionError("must not recompute report")

    monkeypatch.setattr(router, "_dataset", reject)
    monkeypatch.setattr(export, "thumbnail", lambda *args: None)
    monkeypatch.setattr(export, "ensure_headroom", lambda: None)
    result = asyncio.run(router.export_selected(export.SelectedExport(rows=[row()])))
    assert result.status_code == 200


def test_redirect_and_oversized_dimensions_are_rejected(monkeypatch, tmp_path):
    from urllib.error import HTTPError

    with pytest.raises(HTTPError):
        export.NoRedirect().redirect_request(
            export.Request("https://catalog.test/a.png"),
            None,
            302,
            "redirect",
            {},
            "http://127.0.0.1/",
        )
    monkeypatch.setattr(
        export.settings, "vrcatalog_api_url", "https://catalog.test/api"
    )
    monkeypatch.setattr(export, "ensure_headroom", lambda: None)
    data = BytesIO()
    Image.new("RGB", (20, 20), "blue").save(data, format="PNG")

    class Opener:
        def open(self, *args, **kwargs):
            return BytesIO(data.getvalue())

    monkeypatch.setattr(export, "build_opener", lambda *args: Opener())
    monkeypatch.setattr(export, "MAX_IMAGE_PIXELS", 100)
    assert (
        export.thumbnail(
            "https://catalog.test/a.png", tmp_path, 1, time.perf_counter() + 10
        )
        is None
    )
    assert not list(tmp_path.iterdir())
