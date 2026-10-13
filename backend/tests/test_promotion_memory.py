import asyncio
from contextlib import asynccontextmanager
from datetime import date
from io import BytesIO

import httpx
import pytest
from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook

from app.database import get_db
from app.routers import promotion_candidates_report as router
from app.services import product_analytics_runtime as runtime, vrcatalog
from app.services.product_analytics_cache import _bounded_size
from benchmarks import promotion_baseline as before


def legacy(result):
    rows, available = result
    return [
        {key: value for key, value in row.items() if not key.startswith("season")}
        for row in rows
    ], available


from benchmarks.benchmark_promotion_memory import sales_rows, catalog_payload


@pytest.fixture(autouse=True)
def catalog_cache():
    vrcatalog.catalog_info_cache.clear()
    vrcatalog._catalog_info_sizes.clear()
    vrcatalog._catalog_info_bytes = 0
    yield
    vrcatalog.catalog_info_cache.clear()
    vrcatalog._catalog_info_sizes.clear()
    vrcatalog._catalog_info_bytes = 0


def clear_catalog():
    vrcatalog.catalog_info_cache.clear()
    vrcatalog._catalog_info_sizes.clear()
    vrcatalog._catalog_info_bytes = 0


async def no_clients(*args):
    return None


@pytest.mark.parametrize(
    "filters",
    [
        {},
        {"brand": "Бренд"},
        {"manufacturer": "Завод"},
        {"category": "Категория 1"},
        {"reason": "Продажи падают"},
        {"status": "Кандидат"},
        {"brand": "Нет"},
    ],
)
def test_identical_filters_criteria_and_sort(monkeypatch, filters):
    async def sales(*args):
        return list(sales_rows(601))

    for module in (before, router):
        monkeypatch.setattr(module, "_sales", sales, raising=False)
        monkeypatch.setattr(module, "get_sales_filter_clients", no_clients)
    monkeypatch.setattr(
        vrcatalog,
        "_integration_batch",
        lambda payload: catalog_payload(payload["products"]),
    )
    data = router.Request(
        date_from=date(2026, 1, 1), date_to=date(2026, 9, 30), **filters
    )

    async def run():
        original = await before._dataset(data, None)
        clear_catalog()
        optimized = await router._dataset(data, None)
        assert legacy(optimized) == original

    asyncio.run(run())


def test_alias_conflicts_across_5000_boundary_and_compact_warm_cache(monkeypatch):
    async def sales(*args):
        rows = list(sales_rows(5002))
        rows[0]["article"] = "SHARED"
        rows[-1]["article"] = "SHARED"
        rows.append({**rows[-1], "key": "article:shared", "code": None})
        return rows

    for module in (before, router):
        monkeypatch.setattr(module, "_sales", sales, raising=False)
        monkeypatch.setattr(module, "get_sales_filter_clients", no_clients)
    calls = []

    def integration(payload):
        calls.append(len(payload["products"]))
        products = [p for p in payload["products"] if p["code"]]
        return catalog_payload(products)

    monkeypatch.setattr(vrcatalog, "_integration_batch", integration)
    data = router.Request(
        date_from=date(2026, 1, 1), date_to=date(2026, 9, 30), use_seasonality=False
    )

    async def run():
        old = await before._dataset(data, None)
        old_calls = calls.copy()
        clear_catalog()
        calls.clear()
        new = await router._dataset(data, None)
        assert legacy(new) == old and len(new[0]) == 5003
        assert calls == old_calls
        # Both cached and fresh values are projected, never full batch-info.
        compact, available = await router._catalog(await sales())
        assert available
        assert all(
            set(info) == {"category", "brand", "manufacturer", "stock"}
            for info in compact.values()
        )
        assert _bounded_size(compact, 64 * 1024 * 1024) < 5 * 1024 * 1024

    asyncio.run(run())


def test_catalog_unavailable_preserves_old_fallback(monkeypatch):
    async def sales(*args):
        return list(sales_rows(3))

    async def fail(*args, **kwargs):
        raise vrcatalog.VrCatalogError("secret")

    for module in (before, router):
        monkeypatch.setattr(module, "_sales", sales, raising=False)
        monkeypatch.setattr(module, "get_sales_filter_clients", no_clients)
        monkeypatch.setattr(module, "get_catalog_batch_info", fail)
    data = router.Request(
        date_from=date(2026, 1, 1), date_to=date(2026, 9, 30), use_seasonality=False
    )

    async def run():
        assert legacy(await router._dataset(data, None)) == await before._dataset(
            data, None
        )

    asyncio.run(run())


def test_write_only_excel_has_all_rows_and_same_values():
    rows = list(sales_rows(401))
    for row in rows:
        row.update(
            category="Посуда",
            brand="Бренд",
            reasons=["Избыточный запас"],
            status="Кандидат",
            units_change_percent=-60,
            category_change_percent=-60,
            stock=100,
            stock_months=10,
            days_without_sales=60,
        )
    response = router._excel(rows)
    book = load_workbook(BytesIO(response.body), read_only=True)
    try:
        assert book.sheetnames == ["Кандидаты", "История акций"]
        values = list(book["Кандидаты"].values)
        assert len(values) == 402
        assert values[1][:7] == (
            "Товар 0",
            "0000000",
            "A0000000",
            "Посуда",
            "Бренд",
            "Избыточный запас",
            "Кандидат",
        )
        assert values[-1][0] == "Товар 400"
        original = load_workbook(BytesIO(before._excel(rows).body), read_only=True)
        try:
            assert list(original["Кандидаты"].values) == [row[:16] for row in values]
            assert list(original["История акций"].values) == list(
                book["История акций"].values
            )
        finally:
            original.close()
    finally:
        book.close()


def test_excel_memory_rejection_cleans_temporary_xml(monkeypatch):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    initial = set(ALL_TEMP_FILES)

    def reject():
        raise HTTPException(503, "busy")

    monkeypatch.setattr(router, "ensure_headroom", reject)
    with pytest.raises(HTTPException):
        router._excel(list(sales_rows(1)))
    assert set(ALL_TEMP_FILES) == initial


def test_memory_check_in_thread_is_controlled_503(monkeypatch):
    monkeypatch.setattr(runtime, "available_memory", lambda: 1)

    async def run():
        with pytest.raises(HTTPException) as exc:
            await runtime.run_cpu(runtime.ensure_headroom)
        assert exc.value.status_code == 503

    asyncio.run(run())


def application():
    app = FastAPI()
    app.include_router(router.router)

    async def db():
        yield None

    app.dependency_overrides[get_db] = db

    @app.get("/health")
    async def health():
        return {"ok": True}

    return app


@pytest.mark.parametrize("different", [False, True])
def test_whole_report_admission_rejects_parallel_work_and_keeps_health(
    monkeypatch, different
):
    monkeypatch.setattr(runtime, "WAIT_SECONDS", 0.03)

    async def run():
        started = asyncio.Event()
        release = asyncio.Event()
        calls = 0

        async def dataset(*args):
            nonlocal calls
            calls += 1
            started.set()
            await release.wait()
            return [], True

        monkeypatch.setattr(router, "_dataset", dataset)
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=application()), base_url="http://test"
        ) as client:
            body = {"date_from": "2026-01-01", "date_to": "2026-09-30"}
            first = asyncio.create_task(
                client.post("/reports/promotion-candidates", json=body)
            )
            await started.wait()
            assert (await client.get("/health")).status_code == 200
            assert (
                await client.post("/reports/promotion-candidates/history", json=body)
            ).status_code == 200
            rejected = await client.post(
                "/reports/promotion-candidates",
                json={**body, **({"brand": "Другой"} if different else {})},
            )
            assert rejected.status_code == 503 and calls == 1
            release.set()
            assert (await first).status_code == 200

    asyncio.run(run())


def test_cancellation_releases_admission_and_resources(monkeypatch):
    async def run():
        started = asyncio.Event()
        closed = asyncio.Event()

        async def dataset(*args):
            started.set()
            try:
                await asyncio.Event().wait()
            finally:
                closed.set()

        monkeypatch.setattr(router, "_dataset", dataset)
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=application()), base_url="http://test"
        ) as client:
            body = {"date_from": "2026-01-01", "date_to": "2026-09-30"}
            task = asyncio.create_task(
                client.post("/reports/promotion-candidates", json=body)
            )
            await started.wait()
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await task
            assert closed.is_set()
            assert runtime._states[asyncio.get_running_loop()]["active"] == 0

            async def again(*args):
                return [], True

            monkeypatch.setattr(router, "_dataset", again)
            assert (
                await client.post("/reports/promotion-candidates", json=body)
            ).status_code == 200

    asyncio.run(run())


def test_streamed_sales_match_original_sql_aggregates(monkeypatch):
    from test_catalog_directory import database

    data = router.Request(
        date_from=date(2026, 9, 1), date_to=date(2026, 9, 30), use_seasonality=False
    )

    async def run():
        async with database([("001", "A"), ("001", "A"), ("002", "B")]) as (_, factory):
            async with factory() as db:
                old = await before._sales(data, db, None)
                new = await router._sales(data, db, None)
                assert new == old and len(new) == 2
                assert new[0]["units"] == 2

    asyncio.run(run())


def test_sales_cursor_is_closed_on_memory_error(monkeypatch):
    from types import SimpleNamespace

    class Cursor:
        closed = False

        def __aiter__(self):
            async def rows():
                for row in sales_rows(1):
                    yield SimpleNamespace(**row)

            return rows()

        async def close(self):
            self.closed = True

    cursor = Cursor()

    class Database:
        async def stream(self, query):
            assert query.get_execution_options()["yield_per"] == 250
            return cursor

    def reject():
        raise HTTPException(503, "busy")

    monkeypatch.setattr(router, "ensure_headroom", reject)

    async def run():
        with pytest.raises(HTTPException):
            await router._sales(
                router.Request(
                    date_from=date(2026, 1, 1),
                    date_to=date(2026, 9, 30),
                    use_seasonality=False,
                ),
                Database(),
                None,
            )
        assert cursor.closed

    asyncio.run(run())


def test_projection_never_replaces_shared_full_cache(monkeypatch):
    calls = []

    def integration(payload):
        calls.append(True)
        return catalog_payload(payload["products"])

    monkeypatch.setattr(vrcatalog, "_integration_batch", integration)

    async def run():
        products = [{"code": "0000001", "article": "A0000001"}]
        compact = await vrcatalog.get_catalog_batch_info(
            products, project=router._promotion_projection
        )
        full = await vrcatalog.get_catalog_batch_info(products)
        assert set(compact["code:0000001"]) == {
            "category",
            "brand",
            "manufacturer",
            "stock",
        }
        assert "properties" in full["code:0000001"] and "stocks" in full["code:0000001"]
        assert full["code:0000001"]["image_url"] == "https://catalog.test/0000001.jpg"
        assert len(calls) == 1

    asyncio.run(run())


def test_stage_diagnostics_are_numeric_without_filter_values(monkeypatch, caplog):
    async def sales(*args):
        return list(sales_rows(1))

    async def catalog(*args):
        return {}, True

    monkeypatch.setattr(router, "_sales", sales)
    monkeypatch.setattr(router, "_catalog", catalog)
    monkeypatch.setattr(router, "get_sales_filter_clients", no_clients)
    import logging

    with caplog.at_level(logging.INFO, logger=runtime.log.name):
        asyncio.run(
            router._dataset(
                router.Request(
                    date_from=date(2026, 1, 1),
                    date_to=date(2026, 9, 30),
                    brand="PRIVATE_BRAND",
                    manager="PRIVATE_MANAGER",
                ),
                None,
            )
        )
    assert "promotion_sales" in caplog.text and "promotion_classified" in caplog.text
    assert "PRIVATE_" not in caplog.text


def test_memory_admission_never_starts_dataset(monkeypatch):
    monkeypatch.setattr(runtime, "available_memory", lambda: 1)

    async def forbidden(*args):
        pytest.fail("heavy dataset started without RAM")

    monkeypatch.setattr(router, "_dataset", forbidden)

    async def run():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=application()), base_url="http://test"
        ) as client:
            result = await client.post(
                "/reports/promotion-candidates",
                json={"date_from": "2026-01-01", "date_to": "2026-09-30"},
            )
            assert result.status_code == 503
            assert (await client.get("/health")).status_code == 200

    asyncio.run(run())
