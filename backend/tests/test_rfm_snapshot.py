"""PostgreSQL integration tests use only an explicit isolated test database."""
import asyncio
import os
from contextlib import asynccontextmanager
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import event, func, insert, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.config import settings
from app.database import get_db
from app.models import Base, CatalogLookup, CatalogProduct, ImportBatch, RfmSnapshot, Sale, SaleItem
from app.routers import calltrack_integration, rfm_report
from app.services import rfm_snapshot as snapshots
from app.services.catalog_attributes import identity
from app.services.calltrack_insights import client_insights
from app.services.rfm import SegmentSettings, enrich
from app.services.rfm_data import rfm_rows

END = date(2026, 10, 15)


@pytest.mark.parametrize("end,start", [
    (date(2026,10,15),date(2025,10,15)), (date(2026,1,1),date(2025,1,1)),
    (date(2024,2,29),date(2023,3,1)), (date(2025,2,28),date(2024,2,28)),
    (date(2024,3,31),date(2023,3,31))])
def test_period_matches_frontend_set_month(end, start):
    assert snapshots.snapshot_period(end) == (start, end)


@asynccontextmanager
async def database():
    url = os.environ.get("RFM_TEST_DATABASE_URL")
    if not url:
        pytest.skip("RFM_TEST_DATABASE_URL is required for PostgreSQL tests")
    assert url.endswith("/rfm_test"), "Use the isolated rfm_test database"
    schema = "rfm_" + uuid4().hex
    admin = create_async_engine(url)
    async with admin.begin() as conn:
        await conn.execute(text(f'CREATE SCHEMA "{schema}"'))
    engine = create_async_engine(url, connect_args={"server_settings": {"search_path": schema}})
    try:
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        yield engine, async_sessionmaker(engine, expire_on_commit=False)
    finally:
        await engine.dispose()
        async with admin.begin() as conn:
            await conn.execute(text(f'DROP SCHEMA "{schema}" CASCADE'))
        await admin.dispose()


async def seed(sessions):
    async with sessions() as db:
        db.add(ImportBatch(id=1, filename="test", file_size=0, status="completed"))
        await db.flush()
        # Boundary dates, historical first purchase, future exclusion, ties,
        # two departments, same product across multiple sales.
        inputs = [("  АЛЬФА ",date(2020,1,1),"100.00","A"),
                  ("Альфа",date(2025,10,15),"100.01","A"),
                  ("альфа",END,"200.02","B"),
                  ("Альфа",date(2026,10,16),"999.00","A"),
                  ("Бета",END,"300.03","A"),
                  ("Гамма",date(2026,9,1),"900.00","A"),
                  ("Дельта",date(2025,10,14),"999.00","A")]
        for i, (name, day, amount, dept) in enumerate(inputs, 1):
            db.add(Sale(id=i,client=name,sale_date=day,total_amount=Decimal(amount),
                        department=dept,document_number=str(i),fingerprint=str(i),import_id=1,social=False))
        await db.flush()
        for i, sale_id in enumerate([2,2,3,5,6],1):
            db.add(SaleItem(id=i,sale_id=sale_id,name="Product",code="p",quantity=Decimal("0.125")))
        now=datetime.now(timezone.utc)
        db.add(CatalogProduct(product_id=1,code="p",category="Категория",horeca=False,
                              synced_at=now,source_updated_at=now))
        await db.flush()
        db.add(CatalogLookup(identity=identity("p",None),code="p",product_id=1,
                             status="matched",matched_by="code"))
        await db.commit()


def test_persistent_population_products_and_restart(monkeypatch):
    async def scenario():
        async with database() as (engine, sessions):
            await seed(sessions)
            statements=[]
            def record(conn,cursor,statement,parameters,context,executemany): statements.append(statement)
            event.listen(engine.sync_engine,"before_cursor_execute",record)
            async with sessions() as db:
                # Controlled missing state; no population queries.
                with pytest.raises(Exception) as error:
                    await client_insights(db,"Альфа")
                assert error.value.status_code == 503
            assert not any("FROM sales" in sql for sql in statements)
            statements.clear()
            built=await snapshots.rebuild_snapshot(engine,as_of=END)
            assert built["client_count"]==3
            assert not any("sale_items" in sql for sql in statements)
            async with sessions() as db:
                legacy=await rfm_rows(db,date(2025,10,15),END,None)
                legacy,bounds=enrich(legacy,END,SegmentSettings())
                alpha=next(r for r in legacy if r["client_key"]=="альфа")
                assert alpha["first_purchase"]==date(2020,1,1)
                assert alpha["frequency"]==2 and alpha["monetary"]==300.03
                assert alpha["unique_products"]==2  # existing per-sale sum semantics
                assert alpha["units"]==.375
                assert (await snapshots.active_snapshot(db)).boundaries==bounds
            # New pool models a fresh worker. Forbid any population computation.
            await engine.dispose()
            async def forbidden(*a,**kw): raise AssertionError("HTTP attempted rebuild")
            monkeypatch.setattr(snapshots,"rfm_rows",forbidden)
            statements.clear()
            async with sessions() as db:
                response=await client_insights(db,"  аЛьФа ")
                assert response.snapshot.id==built["id"]
                assert response.rfm.rfm_code==alpha["rfm_code"]
                assert response.rfm.tags==alpha["tags"]
                assert response.rfm.monetary==Decimal("300.03")
                assert response.rfm.first_purchase==date(2020,1,1)
                assert response.unique_products==1
                assert response.top_categories[0].category=="Категория"
                assert response.top_categories[0].quantity==Decimal(".375")
                assert [(s.sale_id,s.quantity) for s in response.sale_quantities]==[(2,Decimal(".250")),(3,Decimal(".125"))]
            assert len(statements)==3
            assert not any("array_agg" in sql for sql in statements)
            async with sessions() as db:
                response=await client_insights(db,"Неизвестный")
                assert response.status=="client_not_in_snapshot" and response.rfm is None
                assert response.unique_products==0
    asyncio.run(scenario())


def test_atomic_failure_and_concurrent_builds(monkeypatch):
    async def scenario():
        async with database() as (engine,sessions):
            await seed(sessions)
            first=await snapshots.rebuild_snapshot(engine,as_of=END)
            original=snapshots.client_values
            def fail(*args): raise RuntimeError("injected insertion failure")
            monkeypatch.setattr(snapshots,"client_values",fail)
            with pytest.raises(RuntimeError,match="injected"):
                await snapshots.rebuild_snapshot(engine,as_of=END)
            async with sessions() as db:
                assert (await snapshots.active_snapshot(db)).id==first["id"]
                assert await db.scalar(select(func.count()).select_from(RfmSnapshot))==1
            monkeypatch.setattr(snapshots,"client_values",original)
            entered,release=asyncio.Event(),asyncio.Event()
            original_rows=snapshots.rfm_rows
            calls=0
            async def paused(*a,**kw):
                nonlocal calls
                calls+=1
                entered.set()
                await release.wait()
                return await original_rows(*a,**kw)
            monkeypatch.setattr(snapshots,"rfm_rows",paused)
            one=asyncio.create_task(snapshots.rebuild_snapshot(engine,as_of=END,coalesce=True))
            await asyncio.wait_for(entered.wait(),10)
            two=asyncio.create_task(snapshots.rebuild_snapshot(engine,as_of=END,coalesce=True))
            # Reader sees complete A while B is being built.
            async with sessions() as db:
                snapshot,client=await snapshots.lookup_client(db,"Альфа")
                assert snapshot.id==first["id"] and client.frequency==2
            release.set()
            a,b=await asyncio.wait_for(asyncio.gather(one,two),20)
            assert a["id"]==b["id"]!=first["id"] and calls==1
    asyncio.run(scenario())


def test_http_auth_report_filters_and_export(monkeypatch,caplog):
    async def scenario():
        async with database() as (engine,sessions):
            await seed(sessions)
            await snapshots.rebuild_snapshot(engine,as_of=END)
            app=FastAPI()
            app.include_router(calltrack_integration.router,prefix="/api")
            app.include_router(rfm_report.router,prefix="/api")
            async def dependency():
                async with sessions() as db: yield db
            app.dependency_overrides[get_db]=dependency
            monkeypatch.setattr(settings,"calltrack_integration_token","private-test-token")
            async def metadata(): return {},{}
            monkeypatch.setattr(rfm_report,"get_client_metadata",metadata)
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url="http://test") as client:
                path="/api/integrations/calltrack/client-insights"
                assert (await client.post(path,json={"name":"Альфа"})).status_code==401
                headers={"Authorization":"Bearer private-test-token"}
                response=await client.post(path,json={"name":"Альфа"},headers=headers)
                assert response.status_code==200 and response.json()["status"]=="ready"
                assert "private-test-token" not in response.text and "private-test-token" not in caplog.text
                assert (await client.post(path,json={"name":"  "},headers=headers)).status_code==422
                params={"date_from":"2026-10-15","date_to":"2026-10-15","department":"A"}
                report=await client.get("/api/reports/rfm",params=params)
                assert report.status_code==200
                assert report.json()["clients"]["total"]==1
                assert report.json()["clients"]["items"][0]["client"]=="Бета"
                exported=await client.get("/api/reports/rfm/export",params=params)
                assert exported.status_code==200
                book=load_workbook(BytesIO(exported.content))
                assert book["Клиенты"].max_row==2 and book["Клиенты"]["A2"].value=="Бета"
                assert (await client.get("/api/reports/rfm",params={"date_from":"2026-10-16","date_to":"2026-10-15"})).status_code==422
    asyncio.run(scenario())


@pytest.mark.parametrize("status,expected",[("completed",True),("failed",False)])
def test_import_refresh_only_after_success_and_source_unlock(monkeypatch,status,expected):
    from app.services import imports
    async def scenario():
        async with database() as (engine,sessions):
            await seed(sessions)
            monkeypatch.setattr(imports,"engine",engine)
            monkeypatch.setattr(imports,"SessionLocal",sessions)
            async def process(import_id):
                async with sessions() as db:
                    batch=await db.get(ImportBatch,import_id)
                    batch.status=status
                    await db.commit()
            monkeypatch.setattr(imports,"_process_import_with_source",process)
            calls=[]
            async def refresh(bound_engine):
                async with sessions() as db:
                    assert await db.scalar(select(ImportBatch.status).where(ImportBatch.id==1))=="completed"
                    assert await db.scalar(text("SELECT pg_try_advisory_xact_lock(:key)"),
                                           {"key":snapshots.SOURCE_LOCK})
                calls.append(bound_engine)
            monkeypatch.setattr(snapshots,"rebuild_after_sales_update",refresh)
            await imports.process_import(1)
            assert bool(calls)==expected
    asyncio.run(scenario())


def test_refresh_failure_is_logged_and_preserves_previous(monkeypatch,caplog):
    async def scenario():
        async with database() as (engine,sessions):
            await seed(sessions)
            first=await snapshots.rebuild_snapshot(engine,as_of=END)
            async def fail(*a,**kw): raise RuntimeError("synthetic refresh failure")
            monkeypatch.setattr(snapshots,"rfm_rows",fail)
            assert await snapshots.rebuild_after_sales_update(engine) is None
            async with sessions() as db:
                assert (await snapshots.active_snapshot(db)).id==first["id"]
            assert "synthetic refresh failure" in caplog.text
    asyncio.run(scenario())


def test_lookup_from_fresh_python_process():
    import subprocess
    import sys
    async def scenario():
        async with database() as (engine,sessions):
            await seed(sessions)
            first=await snapshots.rebuild_snapshot(engine,as_of=END)
            async with sessions() as db:
                schema=await db.scalar(text("SELECT current_schema()"))
            code='''
import asyncio, os
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from app.services import rfm_snapshot
from app.services.calltrack_insights import client_insights
async def forbidden(*a,**kw): raise AssertionError("population rebuild in new worker")
rfm_snapshot.rfm_rows=forbidden
async def main():
    engine=create_async_engine(os.environ["RFM_TEST_DATABASE_URL"],connect_args={"server_settings":{"search_path":os.environ["RFM_TEST_SCHEMA"]}})
    async with async_sessionmaker(engine)() as db:
        response=await client_insights(db,"АЛЬФА")
        assert response.status=="ready"
        print(response.snapshot.id)
    await engine.dispose()
asyncio.run(main())
'''
            result=await asyncio.to_thread(subprocess.run,[sys.executable,"-c",code],
                env={**os.environ,"RFM_TEST_SCHEMA":schema},capture_output=True,text=True,timeout=20)
            assert result.returncode==0,result.stderr
            assert result.stdout.strip()==first["id"]
    asyncio.run(scenario())
