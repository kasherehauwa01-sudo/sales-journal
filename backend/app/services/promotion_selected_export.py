"""Bounded export of the already displayed/selected report snapshot."""

import tempfile
import time
from io import BytesIO
from pathlib import Path
from typing import Annotated
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener

from fastapi import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image, UnidentifiedImageError
from pydantic import BaseModel, ConfigDict, Field
from app.config import settings
from app.services.catalog_sync import sync_image_url
from app.services.product_analytics_runtime import ensure_headroom

MAX_IMAGE_BYTES = 2 * 1024 * 1024
MAX_IMAGE_PIXELS = 4_000_000
IMAGE_BUDGET_SECONDS = 15
HEADERS = [
    "Фото",
    "Товар",
    "Код / артикул",
    "Категория",
    "Сезонность",
    "Сезонный приоритет",
    "Причины",
    "30 дней",
    "Пред. 30",
    "Изменение, %",
    "Изменение категории, %",
    "Остаток",
    "Остатки Авиаторов Зал+Склад",
    "Запас, мес.",
    "Без продажи, дней",
    "Статус",
]


class Season(BaseModel):
    status: str = Field(default="Недостаточно данных", max_length=100)
    priority: float = Field(default=0, ge=0, le=100, allow_inf_nan=False)


class SelectedRow(BaseModel):
    model_config = ConfigDict(allow_inf_nan=False)
    key: str = Field(max_length=1000)
    name: str = Field(max_length=1000)
    code: str | None = Field(default=None, max_length=500)
    article: str | None = Field(default=None, max_length=500)
    category: str = Field(default="", max_length=1000)
    image_url: str | None = Field(default=None, max_length=2048)
    seasonality: Season | None = None
    reasons: list[Annotated[str, Field(max_length=1000)]] = Field(
        default_factory=list, max_length=30
    )
    units_30: float = 0
    units_previous_30: float = 0
    units_change_percent: float | None = None
    category_change_percent: float | None = None
    stock: float | None = None
    aviatorov_stock: float | None = None
    stock_months: float | None = None
    days_without_sales: int | None = None
    status: str = Field(default="", max_length=100)


class SelectedExport(BaseModel):
    rows: list[SelectedRow] = Field(min_length=1, max_length=300)


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        raise HTTPError(req.full_url, code, "Redirect blocked", headers, fp)


def trusted_image_url(value):
    url = sync_image_url(value)
    if not url:
        return None
    parsed, source = urlsplit(url), urlsplit(settings.vrcatalog_api_url)
    # Request payload is untrusted. Fetch only the configured catalogue origin,
    # without credentials, redirects or arbitrary third-party destinations.
    if (parsed.scheme, parsed.hostname, parsed.port) != (
        source.scheme,
        source.hostname,
        source.port,
    ):
        return None
    return url


def thumbnail(value, directory, index, deadline):
    url = trusted_image_url(value)
    if not url or time.perf_counter() >= deadline:
        return None
    raw_path = Path(directory) / "source-image"
    target = Path(directory) / f"image-{index}.png"
    try:
        request = Request(
            url,
            headers={"Accept": "image/*", "User-Agent": "SalesJournal-selected-export"},
        )
        timeout = min(3, max(0.1, deadline - time.perf_counter()))
        with build_opener(NoRedirect()).open(
            request, timeout=timeout
        ) as response, raw_path.open("wb") as output:
            size = 0
            while True:
                ensure_headroom()
                if time.perf_counter() >= deadline:
                    return None
                chunk = response.read(64 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_IMAGE_BYTES:
                    return None
                output.write(chunk)
        ensure_headroom()
        with Image.open(raw_path) as picture:
            if picture.width * picture.height > MAX_IMAGE_PIXELS:
                return None
            picture.thumbnail((96, 72))
            if picture.mode not in ("RGB", "RGBA"):
                picture = picture.convert("RGB")
            picture.save(target, format="PNG")
        return target
    except (
        OSError,
        ValueError,
        HTTPError,
        URLError,
        UnidentifiedImageError,
        Image.DecompressionBombError,
    ):
        return None
    finally:
        raw_path.unlink(missing_ok=True)


def safe_cell(sheet, value):
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # Product names must not become Excel formulas.
    return cell


def excel_selected(data):
    ensure_headroom()
    book = Workbook(write_only=True)
    try:
        with tempfile.TemporaryDirectory(prefix="promotion-selected-") as directory:
            sheet = book.create_sheet("Выбранные товары")
            sheet.freeze_panes = "D2"
            sheet.column_dimensions["A"].width = 15
            sheet.column_dimensions["B"].width = 35
            sheet.append([safe_cell(sheet, value) for value in HEADERS])
            deadline = time.perf_counter() + IMAGE_BUDGET_SECONDS
            images = {}
            missing = 0
            for index, row in enumerate(data.rows, 2):
                ensure_headroom()
                if row.image_url not in images:
                    images[row.image_url] = thumbnail(
                        row.image_url, directory, index, deadline
                    )
                path = images[row.image_url]
                if path:
                    sheet.row_dimensions[index].height = 56
                    image = ExcelImage(path)
                    sheet.add_image(image, f"A{index}")
                    photo = ""
                else:
                    photo = sync_image_url(row.image_url) or "Нет фото"
                    missing += 1
                season = row.seasonality or Season()
                values = [
                    photo,
                    row.name,
                    row.code or row.article or "—",
                    row.category,
                    season.status,
                    season.priority,
                    ", ".join(row.reasons),
                    row.units_30,
                    row.units_previous_30,
                    row.units_change_percent,
                    row.category_change_percent,
                    row.stock,
                    row.aviatorov_stock,
                    row.stock_months,
                    row.days_without_sales,
                    row.status,
                ]
                sheet.append([safe_cell(sheet, value) for value in values])
            sheet.auto_filter.ref = f"A1:P{len(data.rows)+1}"
            notes = book.create_sheet("Информация")
            notes.append(
                [
                    "Экспорт выбранных строк из последнего сформированного отчёта; все колонки независимо от настройки видимости."
                ]
            )
            notes.append(
                [
                    f"Строк без встроенного фото: {missing}. Если фото не загрузилось, в столбце сохранена публичная ссылка или «Нет фото»."
                ]
            )
            notes.append(
                [
                    "Фото: последовательные загрузки, до 2 МБ/4 млн пикселей, общий бюджет 15 секунд; сторонние origin и redirect не загружаются."
                ]
            )
            output = BytesIO()
            ensure_headroom()
            book.save(output)
            ensure_headroom()
            return Response(
                output.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": "attachment; filename=promotion-candidates-selected.xlsx"
                },
            )
    finally:
        for sheet in book.worksheets:
            try:
                if not sheet.closed:
                    sheet.close()
                if sheet._writer:
                    sheet._writer.cleanup()
            except (OSError, ValueError):
                pass
        book.close()
