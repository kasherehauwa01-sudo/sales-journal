from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import and_, case, distinct, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models import Sale, SaleItem
from app.services.product_analytics import GROUP_FIELDS, classify, group_rows, merge_periods, previous_period, product_key, summary
from app.services.vrcatalog import VrCatalogError, get_catalog_batch_info

router=APIRouter(prefix="/reports/product-analytics",tags=["Отчеты"])

class AnalyticsRequest(BaseModel):
 date_from:date;date_to:date;compare_from:date|None=None;compare_to:date|None=None;departments:list[str]=Field(default_factory=list);group_by:str="product";brand:str|None=None;manufacturer:str|None=None;category:str|None=None;subcategory:str|None=None;material:str|None=None;article:str|None=None;search:str|None=None

def _dates(data):
 if data.date_from>data.date_to:raise HTTPException(422,"Дата начала не может быть позже даты окончания")
 if bool(data.compare_from)!=bool(data.compare_to):raise HTTPException(422,"Укажите обе даты периода сравнения")
 return (data.compare_from,data.compare_to) if data.compare_from else previous_period(data.date_from,data.date_to)

def _conditions(data,start,end):
 result=[Sale.sale_date>=start,Sale.sale_date<=end]
 if data.departments:result.append(Sale.department.in_(data.departments))
 if data.article:result.append(or_(SaleItem.article.ilike(f"%{data.article.strip()}%"),SaleItem.code.ilike(f"%{data.article.strip()}%")))
 if data.search:result.append(SaleItem.name.ilike(f"%{data.search.strip()}%"))
 return result

async def _period_rows(data,db,start,end):
 revenue=func.coalesce(func.sum(SaleItem.quantity*SaleItem.actual_price),0);units=func.coalesce(func.sum(SaleItem.quantity),0)
 fallback_name=case((and_(or_(SaleItem.article.is_(None),func.trim(SaleItem.article)==""),or_(SaleItem.code.is_(None),func.trim(SaleItem.code)=="")),func.lower(func.trim(SaleItem.name))),else_="")
 q=select(SaleItem.article,SaleItem.code,func.max(SaleItem.name).label("name"),revenue.label("revenue"),units.label("units"),func.count(distinct(Sale.id)).label("checks"),func.min(Sale.sale_date).label("first_sale"),func.max(Sale.sale_date).label("last_sale")).join(Sale,Sale.id==SaleItem.sale_id).where(*_conditions(data,start,end)).group_by(SaleItem.article,SaleItem.code,fallback_name)
 rows=(await db.execute(q)).all();return [{"key":product_key(x.article,x.code,x.name),"article":x.article,"code":x.code,"name":x.name,"revenue":float(x.revenue),"units":float(x.units),"checks":x.checks,"first_sale":x.first_sale,"last_sale":x.last_sale} for x in rows]

async def _catalog(rows):
 products=[{"article":x.get("article"),"code":x.get("code")} for x in rows if x.get("article") or x.get("code")];result={}
 try:
  for start in range(0,len(products),5000):result.update(await get_catalog_batch_info(products[start:start+5000]))
 except VrCatalogError:pass  # отсутствие CatalogVR не скрывает продажи
 return result

async def _dataset(data,db):
 if data.group_by not in GROUP_FIELDS:raise HTTPException(422,"Неизвестная группировка")
 old_from,old_to=_dates(data);current=await _period_rows(data,db,data.date_from,data.date_to);old=await _period_rows(data,db,old_from,old_to);catalog=await _catalog(current+old);rows=merge_periods(current,old,catalog)
 for field in ("brand","manufacturer","category","subcategory","material"):
  value=getattr(data,field)
  if value:rows=[x for x in rows if x[field].casefold()==value.strip().casefold()]
 return group_rows(rows,data.group_by),old_from,old_to

def _sorted(rows,section,sort_by="revenue",limit=100):
 groups=classify(rows) if section!="top" else {}
 source=rows if section=="top" else groups[section]
 if section=="growth":key="revenue_difference" if sort_by!="percent" else "revenue_change";reverse=True
 elif section=="decline":key="revenue_difference";reverse=False
 elif section=="stopped":key="previous_revenue";reverse=True
 else:key=sort_by if sort_by in {"revenue","units","checks"} else "revenue";reverse=True
 return sorted(source,key=lambda x:x.get(key) if x.get(key) is not None else float("-inf"),reverse=reverse)[:limit]

@router.post("/summary")
async def report_summary(data:AnalyticsRequest,db:AsyncSession=Depends(get_db)):
 # KPI всегда считаются по SKU, независимо от выбранной группировки таблицы.
 product_data=data.model_copy(update={"group_by":"product"});rows,old_from,old_to=await _dataset(product_data,db);result=summary(rows)
 # Без фильтров CatalogVR можно получить реальное число уникальных чеков
 # напрямую в PostgreSQL, не суммируя чеки отдельных SKU.
 if not any((data.brand,data.manufacturer,data.category,data.subcategory,data.material)):
  current_checks=await db.scalar(select(func.count(distinct(Sale.id))).join(SaleItem,SaleItem.sale_id==Sale.id).where(*_conditions(data,data.date_from,data.date_to))) or 0
  old_checks=await db.scalar(select(func.count(distinct(Sale.id))).join(SaleItem,SaleItem.sale_id==Sale.id).where(*_conditions(data,old_from,old_to))) or 0
  result["checks"]={"current":current_checks,"previous":old_checks}
 return {"period":{"start":data.date_from,"end":data.date_to},"comparison":{"start":old_from,"end":old_to},"summary":result}

async def _section(data,db,name,sort_by,limit):rows,_,_=await _dataset(data,db);return {"items":_sorted(rows,name,sort_by,limit),"total":len(_sorted(rows,name,sort_by,1000000))}
@router.post("/top")
async def top(data:AnalyticsRequest,sort_by:str="revenue",limit:int=Query(20,ge=10,le=100),db:AsyncSession=Depends(get_db)):return await _section(data,db,"top",sort_by,limit)
@router.post("/growth")
async def growth(data:AnalyticsRequest,sort_by:str="absolute",limit:int=Query(100,le=500),db:AsyncSession=Depends(get_db)):return await _section(data,db,"growth",sort_by,limit)
@router.post("/decline")
async def decline(data:AnalyticsRequest,limit:int=Query(100,le=500),db:AsyncSession=Depends(get_db)):return await _section(data,db,"decline","absolute",limit)
@router.post("/stopped")
async def stopped(data:AnalyticsRequest,limit:int=Query(100,le=500),db:AsyncSession=Depends(get_db)):return await _section(data,db,"stopped","revenue",limit)
@router.post("/new")
async def new(data:AnalyticsRequest,limit:int=Query(100,le=500),db:AsyncSession=Depends(get_db)):return await _section(data,db,"new","revenue",limit)

class DetailRequest(AnalyticsRequest):article_key:str
@router.post("/details")
async def details(data:DetailRequest,group_by:str=Query("day",pattern="^(day|week|month)$"),db:AsyncSession=Depends(get_db)):
 prefix,value=data.article_key.split(":",1);condition=func.lower(func.trim(SaleItem.code if prefix=="code" else SaleItem.article))==value;bucket=func.date_trunc(group_by,Sale.sale_date).label("period");q=select(bucket,func.sum(SaleItem.quantity*SaleItem.actual_price).label("revenue"),func.sum(SaleItem.quantity).label("units"),func.count(distinct(Sale.id)).label("checks")).join(Sale,Sale.id==SaleItem.sale_id).where(*_conditions(data,data.date_from,data.date_to),condition).group_by(bucket).order_by(bucket);points=(await db.execute(q)).all();rows,_,_=await _dataset(data,db);item=next((x for x in rows if x["key"]==data.article_key),None)
 if not item:raise HTTPException(404,"Товар не найден")
 return {"product":item,"group_by":group_by,"points":[{"period":x.period.date(),"revenue":float(x.revenue or 0),"units":float(x.units or 0),"checks":x.checks} for x in points]}

def _sheet(book,title,headers,rows):
 sheet=book.create_sheet(title);sheet.append(headers)
 for row in rows:sheet.append(row)
 sheet.freeze_panes="A2"
 for i,column in enumerate(sheet.columns,1):sheet.column_dimensions[get_column_letter(i)].width=min(45,max(12,max(len(str(x.value or "")) for x in column)+2))

@router.post("/export")
async def export(data:AnalyticsRequest,db:AsyncSession=Depends(get_db)):
 rows,old_from,old_to=await _dataset(data,db);sections={"ТОП":_sorted(rows,"top","revenue",1000000),"Рост":_sorted(rows,"growth","absolute",1000000),"Падение":_sorted(rows,"decline","absolute",1000000),"Перестали продаваться":_sorted(rows,"stopped","revenue",1000000),"Новые товары":_sorted(rows,"new","revenue",1000000)};book=Workbook();book.remove(book.active);s=summary(rows);_sheet(book,"Сводка",["Показатель","Значение"],[[k,str(v)] for k,v in s.items()]+[["Основной период",f"{data.date_from} — {data.date_to}"],["Период сравнения",f"{old_from} — {old_to}"]]);headers=["Артикул","Код","Название","Бренд","Категория","Выручка","Прошлая выручка","Изменение","Изменение, %","Количество","Прошлое количество","Чеков"]
 for title,items in sections.items():_sheet(book,title,headers,[[x.get("article"),x.get("code"),x.get("name"),x.get("brand"),x.get("category"),x.get("revenue"),x.get("previous_revenue"),x.get("revenue_difference"),x.get("revenue_change"),x.get("units"),x.get("previous_units"),x.get("checks")] for x in items]);output=BytesIO();book.save(output);return Response(output.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=product-analytics.xlsx"})
