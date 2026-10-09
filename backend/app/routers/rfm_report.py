from datetime import date, datetime
from io import BytesIO
from math import ceil

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import distinct, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models import Sale, SaleItem
from app.services.clients_vr import (ClientsVrError, cached_client_buyer_types,
    get_buyer_type_clients, get_cached_client_managers, get_client_metadata,
    get_manager_clients, normalize_client_name)
from app.services.rfm import SEGMENTS, SegmentSettings, enrich

router = APIRouter(prefix="/reports/rfm", tags=["Отчеты"])
SORTS = {"client", "segment", "recency_days", "frequency", "monetary", "average_check", "last_purchase"}


async def _rows(db: AsyncSession, start: date, end: date, departments: list[str] | None) -> list[dict]:
    item_stats = (select(
        SaleItem.sale_id.label("sale_id"), func.coalesce(func.sum(SaleItem.quantity), 0).label("units"),
        func.count(distinct(func.coalesce(SaleItem.code, SaleItem.article, SaleItem.name))).label("products"),
    ).group_by(SaleItem.sale_id).subquery())
    client_key = func.lower(func.trim(Sale.client))
    conditions = [Sale.sale_date.between(start, end), Sale.client.is_not(None), func.length(func.trim(Sale.client)) > 0]
    if departments: conditions.append(Sale.department.in_(departments))
    query = (select(
        client_key.label("client_key"), func.min(Sale.client).label("client"),
        func.min(Sale.sale_date).label("period_first"), func.max(Sale.sale_date).label("last_purchase"),
        func.count(Sale.id).label("frequency"), func.sum(Sale.total_amount).label("monetary"),
        func.avg(Sale.total_amount).label("average_check"), func.max(Sale.total_amount).label("max_check"),
        func.min(Sale.total_amount).label("min_check"), func.coalesce(func.sum(item_stats.c.units), 0).label("units"),
        func.coalesce(func.sum(item_stats.c.products), 0).label("unique_products"),
        func.count(distinct(Sale.department)).label("departments_count"),
        func.array_agg(distinct(Sale.sale_date)).label("purchase_dates"),
    ).outerjoin(item_stats, item_stats.c.sale_id == Sale.id).where(*conditions).group_by(client_key))
    current = (await db.execute(query)).mappings().all()
    if not current: return []
    # Первая известная покупка намеренно ищется за всю историю до конца среза.
    history = dict((await db.execute(select(client_key, func.min(Sale.sale_date)).where(
        Sale.sale_date <= end, Sale.client.is_not(None), func.length(func.trim(Sale.client)) > 0
    ).group_by(client_key))).all())
    months = max(1, (end - start).days + 1) / 30.4375
    result = []
    for source in current:
        dates = sorted(source["purchase_dates"] or [])
        gaps = [(b - a).days for a, b in zip(dates, dates[1:])]
        interval = sum(gaps) / len(gaps) if gaps else None
        recency = (end - source["last_purchase"]).days
        result.append({**dict(source), "first_purchase": history[source["client_key"]],
                       "recency_days": recency, "purchases_per_month": source["frequency"] / months,
                       "average_interval": interval, "cycle_overdue": recency / interval if interval and interval > 0 else None,
                       "monetary": float(source["monetary"] or 0), "average_check": float(source["average_check"] or 0),
                       "max_check": float(source["max_check"] or 0), "min_check": float(source["min_check"] or 0),
                       "units": float(source["units"] or 0)})
    return result


async def _dataset(db, date_from, date_to, departments, managers_filter, buyer_types_filter,
                   new_days, lost_days, sleeping_cycle, lost_cycle):
    if date_from > date_to: raise HTTPException(422, "Дата начала не может быть позже даты окончания")
    rows = await _rows(db, date_from, date_to, departments)
    try:
        manager_clients = [await get_manager_clients(value) for value in managers_filter]
        buyer_clients = [await get_buyer_type_clients(value) for value in buyer_types_filter]
    except ClientsVrError as exc:
        raise HTTPException(502, str(exc)) from exc
    allowed_groups = []
    if manager_clients: allowed_groups.append({name.strip().casefold() for group in manager_clients for name in group})
    if buyer_clients: allowed_groups.append({name.strip().casefold() for group in buyer_clients for name in group})
    for allowed in allowed_groups:
        rows = [row for row in rows if row["client_key"].casefold() in allowed]
    rows, boundaries = enrich(rows, date_to, SegmentSettings(new_days, lost_days, sleeping_cycle, lost_cycle))
    managers, buyer_types = get_cached_client_managers(), cached_client_buyer_types() or {}
    try:
        managers, buyer_types = await get_client_metadata()
        clients_vr_available = bool(managers or buyer_types)
    except ClientsVrError:
        clients_vr_available = False
    for row in rows:
        # SQL и ClientsVR могут по-разному обрабатывать регистр и повторные
        # пробелы, поэтому связываем справочники по общему Python-ключу.
        metadata_key = normalize_client_name(row["client"])
        row["manager"] = managers.get(metadata_key); row["buyer_type"] = buyer_types.get(metadata_key)
        row.pop("purchase_dates", None)
    return rows, boundaries, clients_vr_available


def _filtered(rows, segment, r_score, f_score, m_score, rfm_code, search, min_revenue, min_purchases, min_recency, max_recency):
    return [r for r in rows if
        (not segment or r["segment"] == segment) and (not r_score or r["r_score"] == r_score) and
        (not f_score or r["f_score"] == f_score) and (not m_score or r["m_score"] == m_score) and
        (not rfm_code or r["rfm_code"] == rfm_code) and (not search or search.casefold() in r["client"].casefold()) and
        (min_revenue is None or r["monetary"] >= min_revenue) and (min_purchases is None or r["frequency"] >= min_purchases) and
        (min_recency is None or r["recency_days"] >= min_recency) and (max_recency is None or r["recency_days"] <= max_recency)]


@router.get("")
async def report(date_from: date, date_to: date, department: str | None = None, departments: list[str] | None = Query(None),
                 manager: str | None = None, managers: list[str] | None = Query(None), buyer_type: str | None = None,
                 buyer_types: list[str] | None = Query(None), segment: str | None = None, r_score: int | None = Query(None, ge=1, le=5),
                 f_score: int | None = Query(None, ge=1, le=5), m_score: int | None = Query(None, ge=1, le=5),
                 rfm_code: str | None = Query(None, pattern="^[1-5]{3}$"), search: str | None = None,
                 min_revenue: float | None = None, min_purchases: int | None = None, min_recency: int | None = None,
                 max_recency: int | None = None, page: int = Query(1, ge=1), page_size: int = Query(50, ge=10, le=200),
                 sort: str = "monetary", direction: str = Query("desc", pattern="^(asc|desc)$"), new_days: int = Query(30, ge=1),
                 lost_days: int = Query(180, ge=1), sleeping_cycle: float = Query(1.5, ge=1), lost_cycle: float = Query(3, ge=1),
                 db: AsyncSession = Depends(get_db)):
    department_values = departments or ([department] if department else [])
    manager_values = managers or ([manager] if manager else [])
    buyer_type_values = buyer_types or ([buyer_type] if buyer_type else [])
    rows, boundaries, clients_vr_available = await _dataset(db, date_from, date_to, department_values,
        manager_values, buyer_type_values, new_days, lost_days, sleeping_cycle, lost_cycle)
    total_revenue = sum(r["monetary"] for r in rows)
    segments = [{"segment": name, "clients": len(part := [r for r in rows if r["segment"] == name]),
                 "client_share": len(part) / len(rows) if rows else 0, "revenue": sum(r["monetary"] for r in part),
                 "revenue_share": sum(r["monetary"] for r in part) / total_revenue if total_revenue else 0,
                 "average_check": sum(r["monetary"] for r in part) / sum(r["frequency"] for r in part) if part else 0,
                 "average_purchases": sum(r["frequency"] for r in part) / len(part) if part else 0,
                 "average_recency": sum(r["recency_days"] for r in part) / len(part) if part else 0} for name in SEGMENTS]
    selected = _filtered(rows, segment, r_score, f_score, m_score, rfm_code, search, min_revenue, min_purchases, min_recency, max_recency)
    sort = sort if sort in SORTS else "monetary"; selected.sort(key=lambda x: (x[sort] is not None, x[sort]), reverse=direction == "desc")
    offset = (page - 1) * page_size
    counts = {name: sum(r["segment"] == name for r in rows) for name in SEGMENTS}
    inactive = {"Потерянные", "Неактивные"}
    return {"period": {"start": date_from, "end": date_to}, "kpi": {"total_clients": len(rows), "total_revenue": total_revenue,
            "active": sum(r["segment"] not in inactive for r in rows),
            "active_revenue": sum(r["monetary"] for r in rows if r["segment"] not in inactive),
            "sleeping_revenue": sum(r["monetary"] for r in rows if r["segment"] == "Засыпающие"),
            "lost_historical_value": sum(r["monetary"] for r in rows if r["segment"] in inactive), **counts}, "segments": segments,
            "clients": {"items": selected[offset:offset+page_size], "total": len(selected), "page": page,
                        "page_size": page_size, "pages": ceil(len(selected)/page_size) if selected else 0},
            "settings": {"levels": 5, "new_days": new_days, "lost_days": lost_days, "sleeping_cycle": sleeping_cycle,
                         "lost_cycle": lost_cycle, "boundaries": boundaries},
            "clients_vr": {"metadata_available": clients_vr_available}}


@router.get("/clients/{client_key}/history")
async def client_history(client_key: str, date_from: date, date_to: date, db: AsyncSession = Depends(get_db)):
    key = client_key.strip().casefold(); bucket = func.date_trunc("month", Sale.sale_date).label("month")
    rows = (await db.execute(select(bucket, func.sum(Sale.total_amount), func.count(Sale.id)).where(
        Sale.sale_date.between(date_from, date_to), func.lower(func.trim(Sale.client)) == key).group_by(bucket).order_by(bucket))).all()
    return [{"month": row[0].date(), "revenue": float(row[1] or 0), "purchases": row[2]} for row in rows]


@router.get("/export")
async def export(date_from: date, date_to: date, department: str | None = None, departments: list[str] | None = Query(None),
                 manager: str | None = None, managers: list[str] | None = Query(None), buyer_type: str | None = None,
                 buyer_types: list[str] | None = Query(None), segment: str | None = None, r_score: int | None = Query(None, ge=1, le=5),
                 f_score: int | None = Query(None, ge=1, le=5), m_score: int | None = Query(None, ge=1, le=5),
                 rfm_code: str | None = Query(None, pattern="^[1-5]{3}$"), search: str | None = None,
                 min_revenue: float | None = None, min_purchases: int | None = None, min_recency: int | None = None,
                 max_recency: int | None = None, sort: str = "monetary", direction: str = Query("desc", pattern="^(asc|desc)$"),
                 new_days: int = Query(30, ge=1), lost_days: int = Query(180, ge=1), sleeping_cycle: float = Query(1.5, ge=1),
                 lost_cycle: float = Query(3, ge=1), full: bool = True, db: AsyncSession = Depends(get_db)):
    rows, boundaries, _ = await _dataset(db, date_from, date_to, departments or ([department] if department else []),
        managers or ([manager] if manager else []), buyer_types or ([buyer_type] if buyer_type else []),
        new_days, lost_days, sleeping_cycle, lost_cycle)
    rows = _filtered(rows, segment, r_score, f_score, m_score, rfm_code, search, min_revenue, min_purchases, min_recency, max_recency)
    sort = sort if sort in SORTS else "monetary"; rows.sort(key=lambda row: (row[sort] is not None, row[sort]), reverse=direction == "desc")
    wb = Workbook(); ws = wb.active; ws.title = "Клиенты"
    columns = [("Клиент","client"),("Сегмент","segment"),("Теги","tags"),("RFM","rfm_code"),("Последняя покупка","last_purchase"),
               ("Дней без покупки","recency_days"),("Покупок","frequency"),("Выручка","monetary"),("Средний чек","average_check"),
               ("Менеджер","manager"),("Тип покупателя","buyer_type")]
    ws.append([c[0] for c in columns]); ws[1][0].font = Font(bold=True)
    for row in rows: ws.append([", ".join(row[k]) if isinstance(row.get(k), list) else row.get(k) for _, k in columns])
    for cell in ws["D"][1:]: cell.number_format = "@"
    for col in ("H","I"):
        for cell in ws[col][1:]: cell.number_format = '#,##0.00'
    if full:
        seg = wb.create_sheet("Сегменты"); seg.append(["Сегмент","Клиенты","Выручка"])
        for name in SEGMENTS:
            part=[r for r in rows if r["segment"]==name]; seg.append([name,len(part),sum(r["monetary"] for r in part)])
        cfg=wb.create_sheet("Настройки RFM"); cfg.append(["Метод","Квинтили уникальных значений без разрыва совпадений"]); cfg.append(["Границы",str(boundaries)])
    data=BytesIO(); wb.save(data); data.seek(0)
    filename=f"rfm-analysis-{datetime.now().date().isoformat()}.xlsx"
    return StreamingResponse(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})
