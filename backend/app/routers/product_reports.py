from datetime import date,timedelta
from io import BytesIO
from fastapi import APIRouter,Depends,HTTPException,Query,Response
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel,Field
from sqlalchemy import delete,distinct,func,or_,select
from sqlalchemy.ext.asyncio import AsyncSession
from app.database import get_db
from app.models import ProductReportSet,Sale,SaleItem
from app.services.clients_vr import ClientsVrError,get_client_managers,get_manager_clients
from app.services.product_report_utils import normalize_identifier as _norm,percent_change as _change,previous_period as _previous,product_key as _product_key
from app.services.vrcatalog import VrCatalogError,get_product_filter_options,get_product_filters,search_catalog_products

router=APIRouter(prefix="/reports/product-sales",tags=["Отчеты"])

class ProductRef(BaseModel):article:str|None=None;code:str|None=None;name:str
class ReportRequest(BaseModel):
 date_from:date;date_to:date;manager:str|None=None;departments:list[str]=Field(default_factory=list);products:list[ProductRef]=Field(default_factory=list);compare:bool=False
class DetailRequest(ReportRequest):product:ProductRef
class SetIn(BaseModel):name:str=Field(min_length=1,max_length=255);products:list[ProductRef]
class CatalogSearch(BaseModel):filters:dict[str,list[str]]=Field(default_factory=dict);search:str="";page:int=Field(1,ge=1);page_size:int=Field(50,ge=1,le=500)

def _product_condition(products):
 codes={_norm(x.code) for x in products if _norm(x.code)};articles={_norm(x.article) for x in products if not _norm(x.code) and _norm(x.article)};names={_norm(x.name) for x in products if not _norm(x.code) and not _norm(x.article)};parts=[]
 if codes:parts.append(func.lower(func.trim(SaleItem.code)).in_(codes))
 if articles:parts.append(func.lower(func.trim(SaleItem.article)).in_(articles))
 if names:parts.append(func.lower(func.trim(SaleItem.name)).in_(names))
 return or_(*parts) if parts else None
async def _conditions(data:ReportRequest,start=None,end=None):
 conditions=[Sale.sale_date>=(start or data.date_from),Sale.sale_date<=(end or data.date_to)]
 if data.departments:conditions.append(Sale.department.in_(data.departments))
 product_condition=_product_condition(data.products)
 if product_condition is None:conditions.append(False)
 else:conditions.append(product_condition)
 if data.manager:
  try:clients=await get_manager_clients(data.manager)
  except ClientsVrError as exc:raise HTTPException(502,str(exc)) from exc
  normalized=[_norm(x) for x in clients];conditions.append(func.lower(func.trim(Sale.client)).in_(normalized) if normalized else False)
 return conditions

@router.get("/products/search")
async def search_products(search:str="",page:int=Query(1,ge=1),page_size:int=Query(50,ge=1,le=200),db:AsyncSession=Depends(get_db)):
 base=select(SaleItem.article,SaleItem.code,SaleItem.name).group_by(SaleItem.article,SaleItem.code,SaleItem.name)
 if search.strip():
  value=f"%{search.strip()}%";base=base.where(or_(SaleItem.article.ilike(value),SaleItem.code.ilike(value),SaleItem.name.ilike(value)))
 total=await db.scalar(select(func.count()).select_from(base.subquery())) or 0
 rows=(await db.execute(base.order_by(SaleItem.name).offset((page-1)*page_size).limit(page_size))).all()
 return {"items":[{"key":_product_key(*row),"article":row[0],"code":row[1],"name":row[2]} for row in rows],"total":total,"page":page,"page_size":page_size,"pages":max(1,(total+page_size-1)//page_size)}

async def _summary(data,db,start=None,end=None):
 c=await _conditions(data,start,end);revenue=func.coalesce(func.sum(SaleItem.quantity*SaleItem.actual_price),0);units=func.coalesce(func.sum(SaleItem.quantity),0);base=func.coalesce(func.sum(SaleItem.quantity*SaleItem.base_price),0)
 row=(await db.execute(select(revenue,units,func.count(distinct(Sale.id)),func.count(distinct(Sale.client)),base).join(Sale,Sale.id==SaleItem.sale_id).where(*c))).one();rev=float(row[0]);qty=float(row[1]);checks=row[2];discount=float(row[4])-rev
 return {"revenue":rev,"units":qty,"checks":checks,"clients":row[3],"average_price":rev/qty if qty else 0,"items_per_check":qty/checks if checks else 0,"discount_amount":discount,"average_discount":discount/float(row[4])*100 if row[4] else 0}
@router.post("/summary")
async def summary(data:ReportRequest,db:AsyncSession=Depends(get_db)):
 current=await _summary(data,db);previous=await _summary(data,db,*_previous(data.date_from,data.date_to)) if data.compare else None
 return {"current":current,"previous":previous,"changes":{key:{"absolute":current[key]-previous[key],"percent":_change(current[key],previous[key])} for key in current} if previous else None}

async def _products(data,db,start=None,end=None):
 c=await _conditions(data,start,end);revenue=func.sum(SaleItem.quantity*SaleItem.actual_price);units=func.sum(SaleItem.quantity);base=func.sum(SaleItem.quantity*SaleItem.base_price)
 q=select(SaleItem.article,SaleItem.code,SaleItem.name,units.label("units"),revenue.label("revenue"),func.count(distinct(Sale.id)).label("checks"),func.count(distinct(Sale.client)).label("clients"),(revenue/func.nullif(units,0)).label("average_price"),(base-revenue).label("discount_amount"),((base-revenue)/func.nullif(base,0)*100).label("average_discount"),func.max(Sale.sale_date).label("last_sale")).join(Sale,Sale.id==SaleItem.sale_id).where(*c).group_by(SaleItem.article,SaleItem.code,SaleItem.name)
 rows=(await db.execute(q)).all();total=sum(float(x.revenue or 0) for x in rows)
 return [{**dict(x._mapping),"key":_product_key(x.article,x.code,x.name),"units":float(x.units or 0),"revenue":float(x.revenue or 0),"average_price":float(x.average_price or 0),"discount_amount":float(x.discount_amount or 0),"average_discount":float(x.average_discount or 0),"revenue_share":float(x.revenue or 0)/total*100 if total else 0} for x in rows]
@router.post("/products")
async def products(data:ReportRequest,sort_by:str="revenue",sort_dir:str="desc",page:int=Query(1,ge=1),page_size:int=Query(100,ge=1,le=500),db:AsyncSession=Depends(get_db)):
 current=await _products(data,db);previous=await _products(data,db,*_previous(data.date_from,data.date_to)) if data.compare else [];previous_by={x["key"]:x for x in previous}
 for item in current:
  old=previous_by.get(item["key"],{});item["previous_revenue"]=old.get("revenue",0);item["previous_units"]=old.get("units",0);item["revenue_change_percent"]=_change(item["revenue"],item["previous_revenue"]);item["units_change_percent"]=_change(item["units"],item["previous_units"])
 allowed={"revenue","units","checks","clients","average_price","discount_amount","last_sale"};key=sort_by if sort_by in allowed else "revenue";current.sort(key=lambda x:(x[key] is not None,x[key]),reverse=sort_dir=="desc")
 start=(page-1)*page_size;return {"items":current[start:start+page_size],"total":len(current),"page":page,"pages":max(1,(len(current)+page_size-1)//page_size)}

@router.post("/dynamics")
async def dynamics(data:ReportRequest,group_by:str=Query("day",pattern="^(day|week|month)$"),db:AsyncSession=Depends(get_db)):
 c=await _conditions(data);period=func.date_trunc(group_by,Sale.sale_date).label("period");q=select(period,func.sum(SaleItem.quantity*SaleItem.actual_price).label("revenue"),func.sum(SaleItem.quantity).label("units")).join(Sale,Sale.id==SaleItem.sale_id).where(*c).group_by(period).order_by(period)
 return [{"period":x.period.date(),"revenue":float(x.revenue or 0),"units":float(x.units or 0)} for x in (await db.execute(q))]

@router.post("/departments")
async def departments(data:ReportRequest,db:AsyncSession=Depends(get_db)):
 c=await _conditions(data);q=select(Sale.department,func.sum(SaleItem.quantity*SaleItem.actual_price),func.sum(SaleItem.quantity),func.count(distinct(Sale.id)),func.count(distinct(Sale.client))).join(Sale,Sale.id==SaleItem.sale_id).where(*c).group_by(Sale.department).order_by(func.sum(SaleItem.quantity*SaleItem.actual_price).desc())
 return [{"department":x[0],"revenue":float(x[1] or 0),"units":float(x[2] or 0),"checks":x[3],"clients":x[4]} for x in (await db.execute(q))]

@router.post("/managers")
async def managers(data:ReportRequest,db:AsyncSession=Depends(get_db)):
 if data.manager:return []
 c=await _conditions(data);q=select(Sale.client,func.sum(SaleItem.quantity*SaleItem.actual_price),func.sum(SaleItem.quantity),func.count(distinct(Sale.id))).join(Sale,Sale.id==SaleItem.sale_id).where(*c).group_by(Sale.client);rows=(await db.execute(q)).all()
 try:mapping=await get_client_managers()
 except ClientsVrError as exc:raise HTTPException(502,str(exc)) from exc
 totals={}
 for client,revenue,units,checks in rows:
  manager=mapping.get(_norm(client),"Нет менеджера");item=totals.setdefault(manager,{"manager":manager,"revenue":0.0,"units":0.0,"checks":0,"clients":0});item["revenue"]+=float(revenue or 0);item["units"]+=float(units or 0);item["checks"]+=checks;item["clients"]+=1
 return sorted(totals.values(),key=lambda x:-x["revenue"])

@router.post("/details")
async def details(data:DetailRequest,page:int=Query(1,ge=1),page_size:int=Query(50,ge=1,le=200),db:AsyncSession=Depends(get_db)):
 scoped=ReportRequest(**data.model_dump(exclude={"product","products"}),products=[data.product]);c=await _conditions(scoped);q=select(Sale.id,Sale.sale_date,Sale.document_number,Sale.client,Sale.department,SaleItem.quantity,(SaleItem.quantity*SaleItem.actual_price).label("amount"),SaleItem.actual_price,((SaleItem.base_price-SaleItem.actual_price)/func.nullif(SaleItem.base_price,0)*100).label("discount")).join(Sale,Sale.id==SaleItem.sale_id).where(*c).order_by(Sale.sale_date.desc());total=await db.scalar(select(func.count()).select_from(q.subquery())) or 0;rows=(await db.execute(q.offset((page-1)*page_size).limit(page_size))).all()
 try:mapping=await get_client_managers()
 except ClientsVrError as exc:raise HTTPException(502,str(exc)) from exc
 return {"items":[{**dict(x._mapping),"manager":mapping.get(_norm(x.client),"Нет менеджера"),"quantity":float(x.quantity),"amount":float(x.amount or 0),"actual_price":float(x.actual_price or 0),"discount":float(x.discount or 0)} for x in rows],"total":total,"page":page,"pages":max(1,(total+page_size-1)//page_size)}

@router.get("/sets")
async def sets(db:AsyncSession=Depends(get_db)):
 rows=(await db.scalars(select(ProductReportSet).order_by(ProductReportSet.name))).all();return [{"id":row.id,"name":row.name,"products":[{**item,"key":_product_key(item.get("article"),item.get("code"),item.get("name"))} for item in row.products]} for row in rows]
@router.post("/sets")
async def create_set(data:SetIn,db:AsyncSession=Depends(get_db)):
 row=ProductReportSet(name=data.name.strip(),products=[x.model_dump() for x in data.products]);db.add(row)
 try:await db.commit();await db.refresh(row)
 except Exception as exc:await db.rollback();raise HTTPException(409,"Набор с таким названием уже существует") from exc
 return row
@router.put("/sets/{set_id}")
async def update_set(set_id:int,data:SetIn,db:AsyncSession=Depends(get_db)):
 row=await db.get(ProductReportSet,set_id)
 if not row:raise HTTPException(404,"Набор не найден")
 row.name=data.name.strip();row.products=[x.model_dump() for x in data.products];await db.commit();await db.refresh(row);return row
@router.delete("/sets/{set_id}",status_code=204)
async def delete_set(set_id:int,db:AsyncSession=Depends(get_db)):await db.execute(delete(ProductReportSet).where(ProductReportSet.id==set_id));await db.commit();return Response(status_code=204)

@router.get("/catalog/filters")
async def catalog_filters():
 try:return await get_product_filters()
 except VrCatalogError as exc:raise HTTPException(502,str(exc)) from exc

@router.get("/catalog/filters/{filter_key}/options")
async def catalog_filter_options(filter_key:str,search:str="",page:int=Query(1,ge=1),page_size:int=Query(100,ge=1,le=500)):
 try:return await get_product_filter_options(filter_key,search=search,page=page,page_size=page_size)
 except VrCatalogError as exc:raise HTTPException(502,str(exc)) from exc

@router.post("/catalog/products/search")
async def catalog_products_search(data:CatalogSearch):
 try:return await search_catalog_products(filters=data.filters,search=data.search,page=data.page,page_size=data.page_size)
 except VrCatalogError as exc:raise HTTPException(502,str(exc)) from exc

def _sheet(book,title,headers,rows):
 sheet=book.create_sheet(title);sheet.append(headers)
 for row in rows:sheet.append(list(row))
 sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
 for index,column in enumerate(sheet.columns,1):sheet.column_dimensions[get_column_letter(index)].width=min(40,max(12,max(len(str(cell.value or "")) for cell in column)+2))
@router.post("/export")
async def export(data:ReportRequest,db:AsyncSession=Depends(get_db)):
 summary_data=await _summary(data,db);product_data=await _products(data,db);department_data=await departments(data,db);manager_data=await managers(data,db);c=await _conditions(data);detail_q=select(Sale.sale_date,Sale.document_number,Sale.client,Sale.department,SaleItem.quantity,(SaleItem.quantity*SaleItem.actual_price).label("amount"),SaleItem.actual_price,((SaleItem.base_price-SaleItem.actual_price)/func.nullif(SaleItem.base_price,0)*100).label("discount")).join(Sale,Sale.id==SaleItem.sale_id).where(*c).order_by(Sale.sale_date.desc());detail_source=(await db.execute(detail_q)).all()
 try:client_managers=await get_client_managers()
 except ClientsVrError as exc:raise HTTPException(502,str(exc)) from exc
 details_rows=[{**dict(x._mapping),"manager":client_managers.get(_norm(x.client),"Нет менеджера"),"quantity":float(x.quantity),"amount":float(x.amount or 0),"actual_price":float(x.actual_price or 0),"discount":float(x.discount or 0)} for x in detail_source]
 book=Workbook();book.remove(book.active);_sheet(book,"Итоги",["Показатель","Значение"],summary_data.items());_sheet(book,"Товары",["Артикул","Код","Наименование","Продано","Продажи","Чеков","Клиентов","Средняя цена","Скидка","Средняя скидка","Последняя продажа","Доля"],[(x["article"],x["code"],x["name"],x["units"],x["revenue"],x["checks"],x["clients"],x["average_price"],x["discount_amount"],x["average_discount"],x["last_sale"],x["revenue_share"]) for x in product_data]);_sheet(book,"Менеджеры",["Менеджер","Продажи","Продано","Чеков","Клиентов"],[(x["manager"],x["revenue"],x["units"],x["checks"],x["clients"]) for x in manager_data]);_sheet(book,"Подразделения",["Подразделение","Продажи","Продано","Чеков","Клиентов"],[(x["department"],x["revenue"],x["units"],x["checks"],x["clients"]) for x in department_data]);_sheet(book,"Детализация продаж",["Дата","Документ","Клиент","Менеджер","Подразделение","Количество","Сумма","Цена","Скидка"],[(x["sale_date"],x["document_number"],x["client"],x["manager"],x["department"],x["quantity"],x["amount"],x["actual_price"],x["discount"]) for x in details_rows]);output=BytesIO();book.save(output)
 return Response(output.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=product-sales.xlsx"})
