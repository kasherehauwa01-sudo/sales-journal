from datetime import date, datetime
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import String, and_, column, false, func, or_, select, table, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.database import get_db
from app.models import Sale, SaleItem
from app.services.category_analytics import (
    NOT_FOUND,
    UNCATEGORIZED,
    aggregate_categories,
    categorized_only,
    chart_structure,
    comparable_period,
    report_summary,
)
from app.services.report_periods import previous_year_period
from app.services.sales_dynamics_report import previous_period
from app.services.clients_vr import ClientsVrError
from app.services.sales_client_filters import get_sales_filter_clients
from app.services.sales_dynamics_report import default_grouping, effective_period
from app.services.vrcatalog import (
    CATALOG_CATEGORY_MAP_LIMIT,
    VrCatalogError,
    get_catalog_category_map,
)

router = APIRouter(prefix="/reports/category-analytics", tags=["Отчеты"])


def _sale_conditions(start: date, end: date, stores: list[str], clients: list[str] | None):
    result = [Sale.sale_date >= start, Sale.sale_date <= end]
    if stores: result.append(Sale.department.in_(stores))
    if clients is not None:
        names = [name.strip().casefold() for name in clients if name.strip()]
        result.append(func.lower(func.trim(Sale.client)).in_(names) if names else false())
    return result


async def _clients(db: AsyncSession, manager: str | None, buyer_type: str | None):
    try: return await get_sales_filter_clients(db, manager, buyer_type)
    except ClientsVrError as exc: raise HTTPException(502, str(exc)) from exc


async def _product_keys(
    db: AsyncSession,
    start: date,
    end: date,
    previous_start: date,
    previous_end: date,
    stores: list[str],
    clients: list[str] | None,
    additional_period=None,
):
    period = or_(
        and_(Sale.sale_date >= start, Sale.sale_date <= end),
        and_(Sale.sale_date >= previous_start, Sale.sale_date <= previous_end),
    )
    if additional_period:
        period = or_(
            period,
            and_(
                Sale.sale_date >= additional_period[0],
                Sale.sale_date <= additional_period[1],
            ),
        )
    conditions = [period]
    if stores:
        conditions.append(Sale.department.in_(stores))
    if clients is not None:
        names = [name.strip().casefold() for name in clients if name.strip()]
        conditions.append(
            func.lower(func.trim(Sale.client)).in_(names) if names else false()
        )
    rows = (
        await db.execute(
            select(
                SaleItem.code, SaleItem.article, func.max(SaleItem.name).label("name")
            )
            .join(Sale)
            .where(*conditions)
            .group_by(SaleItem.code, SaleItem.article)
        )
    ).all()
    return [
        {"code": row.code, "article": row.article, "name": row.name} for row in rows
    ]


async def _catalog_map(products: list[dict]):
    catalog = {}
    try:
        # Последовательные пакеты не создают пиковую нагрузку на CatalogVR.
        for offset in range(0, len(products), CATALOG_CATEGORY_MAP_LIMIT):
            catalog.update(await get_catalog_category_map(products[offset:offset + CATALOG_CATEGORY_MAP_LIMIT]))
    except VrCatalogError as exc:
        raise HTTPException(502, "Не удалось получить категории товаров из CatalogVR. Попробуйте повторить позже.") from exc
    mapped = []; seen = set()
    for product in products:
        code = str(product.get("code") or "").strip().casefold() or None
        article = str(product.get("article") or "").strip().casefold() or None
        stable_key = ("code", code) if code else ("article", article) if article else None
        if stable_key and stable_key in seen: continue
        if stable_key: seen.add(stable_key)
        code_key = f"code:{code}" if code else None; article_key = f"article:{article}" if article else None
        info = catalog.get(code_key) if code_key else None
        if info is None and article_key: info = catalog.get(article_key)
        if info is None: category = NOT_FOUND
        else: category = str(info.get("category") or "").strip() or UNCATEGORIZED
        mapped.append((code, article, category))
    return mapped, catalog


async def _mapping_table(db: AsyncSession, mapping: list[tuple[str | None, str | None, str]]):
    """Загружает компактную карту CatalogVR во временную таблицу без огромного VALUES."""
    name = f"category_map_{uuid4().hex}"
    await db.execute(text(f"CREATE TEMP TABLE {name} (code text, article text, category text NOT NULL) ON COMMIT DROP"))
    rows = [{"code": code, "article": article, "category": category} for code, article, category in mapping]
    if rows:
        # Список параметров включает executemany: размер одного SQL-запроса не
        # растёт вместе с ассортиментом и не упирается в лимит параметров asyncpg.
        await db.execute(text(f"INSERT INTO {name} (code, article, category) VALUES (:code, :article, :category)"), rows)
        await db.execute(text(f"CREATE INDEX ON {name} (code)")); await db.execute(text(f"CREATE INDEX ON {name} (article)"))
        await db.execute(text(f"ANALYZE {name}"))
    return table(name, column("code", String), column("article", String), column("category", String))


def _mapped_items(category_map):
    code = func.lower(func.trim(SaleItem.code)); article = func.lower(func.trim(SaleItem.article))
    join = or_(and_(category_map.c.code.is_not(None), code == category_map.c.code), and_(or_(SaleItem.code.is_(None), func.trim(SaleItem.code) == ""), category_map.c.code.is_(None), article == category_map.c.article))
    category = func.coalesce(category_map.c.category, UNCATEGORIZED).label("category")
    return category_map, join, category


async def _category_rows(db, start, end, stores, clients, category_map):
    category_map, join, category = _mapped_items(category_map)
    query = select(category, func.coalesce(func.sum(SaleItem.quantity * SaleItem.actual_price), 0).label("revenue"), func.coalesce(func.sum(SaleItem.quantity), 0).label("units"), func.count(func.distinct(Sale.id)).label("checks")).select_from(SaleItem).join(Sale).outerjoin(category_map, join).where(*_sale_conditions(start, end, stores, clients)).group_by(category)
    rows = (await db.execute(query)).all()
    return [{"category": row.category, "revenue": float(row.revenue), "units": float(row.units), "checks": int(row.checks)} for row in rows]


async def _scope_checks(db, start, end, stores, clients):
    query = select(func.count(func.distinct(Sale.id))).select_from(SaleItem).join(Sale).where(*_sale_conditions(start, end, stores, clients))
    return int(await db.scalar(query) or 0)


async def _scope(
    date_from,
    date_to,
    period_kind,
    stores,
    manager,
    buyer_type,
    category_filter,
    db,
    comparison_mode="previous_period",
    include_year=False,
):
    if date_from > date_to:
        raise HTTPException(422, "Дата начала не может быть позже даты окончания")
    today = datetime.now(ZoneInfo(settings.autoload_timezone)).date()
    start, end, warning = effective_period(date_from, date_to, period_kind, today)
    if start > end:
        raise HTTPException(422, warning or "В выбранном периоде пока нет данных")
    previous_start, previous_end = (
        previous_year_period(start, end)
        if comparison_mode == "previous_year"
        else previous_period(start, end, period_kind)
    )
    year_start, year_end = previous_year_period(start, end)
    clients = await _clients(db, manager, buyer_type)
    products = await _product_keys(
        db,
        start,
        end,
        previous_start,
        previous_end,
        stores,
        clients,
        (year_start, year_end) if include_year else None,
    )
    mapping, catalog = await _catalog_map(products)
    category_map = await _mapping_table(db, mapping)
    current = await _category_rows(db, start, end, stores, clients, category_map)
    previous = await _category_rows(
        db, previous_start, previous_end, stores, clients, category_map
    )
    rows = categorized_only(aggregate_categories(current, previous))
    if category_filter:
        rows = [
            row
            for row in rows
            if row["category"].casefold() == category_filter.strip().casefold()
        ]
    if category_filter:
        current_checks = next(
            (
                row["checks"]
                for row in current
                if row["category"].casefold() == category_filter.strip().casefold()
            ),
            0,
        )
        previous_checks = next(
            (
                row["checks"]
                for row in previous
                if row["category"].casefold() == category_filter.strip().casefold()
            ),
            0,
        )
    else:
        current_checks = await _scope_checks(db, start, end, stores, clients)
        previous_checks = await _scope_checks(
            db, previous_start, previous_end, stores, clients
        )
    base = (
        start,
        end,
        previous_start,
        previous_end,
        warning,
        clients,
        category_map,
        catalog,
        rows,
        current_checks,
        previous_checks,
    )
    if include_year:
        year = await _category_rows(
            db, year_start, year_end, stores, clients, category_map
        )
        year_rows = categorized_only(aggregate_categories(current, year))
        year_checks = await _scope_checks(db, year_start, year_end, stores, clients)
        return (*base, year_rows, year_checks, year_start, year_end)
    return base


@router.get("")
async def report(
    date_from: date,
    date_to: date,
    period_kind: str = "custom",
    stores: list[str] = Query(default=[]),
    manager: str | None = None,
    buyer_type: str | None = None,
    category: str | None = None,
    db: AsyncSession = Depends(get_db),
):
    (
        start,
        end,
        old_start,
        old_end,
        warning,
        _,
        _,
        _,
        all_rows,
        checks,
        old_checks,
        year_rows,
        year_checks,
        year_start,
        year_end,
    ) = await _scope(
        date_from,
        date_to,
        period_kind,
        stores,
        manager,
        buyer_type,
        None,
        db,
        include_year=True,
    )
    year_rows = [
        row
        for row in year_rows
        if not category or row["category"].casefold() == category.strip().casefold()
    ]
    if category:
        year_checks = year_rows[0]["previous_checks"] if year_rows else 0
    options = sorted(row["category"] for row in all_rows)
    rows = [
        row
        for row in all_rows
        if not category or row["category"].casefold() == category.strip().casefold()
    ]
    if category:
        checks = rows[0]["current_checks"] if rows else 0
        old_checks = rows[0]["previous_checks"] if rows else 0
    public_rows = [
        {key: value for key, value in row.items() if key != "products"} for row in rows
    ]
    return {
        "period": {"start": start, "end": end},
        "previous_period": {"start": old_start, "end": old_end},
        "warning": warning,
        "metrics": report_summary(rows, checks, old_checks),
        "year_comparison": {
            "period": {"start": year_start, "end": year_end},
            "metrics": report_summary(year_rows, checks, year_checks),
            "growth_drivers": sorted(
                (row for row in year_rows if row["revenue_change"] > 0),
                key=lambda row: row["revenue_change"],
                reverse=True,
            )[:5],
            "decline_drivers": sorted(
                (row for row in year_rows if row["revenue_change"] < 0),
                key=lambda row: row["revenue_change"],
            )[:5],
        },
        "categories": public_rows,
        "category_options": options,
        "structure": chart_structure(rows),
        "growth_drivers": sorted(
            (row for row in public_rows if row["revenue_change"] > 0),
            key=lambda row: row["revenue_change"],
            reverse=True,
        )[:5],
        "decline_drivers": sorted(
            (row for row in public_rows if row["revenue_change"] < 0),
            key=lambda row: row["revenue_change"],
        )[:5],
    }


async def _product_details(db, start, end, old_start, old_end, stores, clients, category_map, category_name):
    category_map, join, category = _mapped_items(category_map)
    async def period_rows(period_start, period_end):
        conditions = _sale_conditions(period_start, period_end, stores, clients)
        if category_name: conditions.append(category == category_name)
        rows = (await db.execute(select(category, SaleItem.article, SaleItem.code, func.max(SaleItem.name).label("name"), func.sum(SaleItem.quantity * SaleItem.actual_price).label("revenue"), func.sum(SaleItem.quantity).label("units"), func.count(func.distinct(Sale.id)).label("checks")).select_from(SaleItem).join(Sale).outerjoin(category_map, join).where(*conditions).group_by(category, SaleItem.article, SaleItem.code))).all()
        return {(row.category, row.code or row.article or row.name): row for row in rows}
    current, previous = await period_rows(start, end), await period_rows(old_start, old_end)
    total = sum(float(row.revenue or 0) for row in current.values())
    result = []
    for key in current.keys() | previous.keys():
        now, old = current.get(key), previous.get(key); revenue = float(now.revenue or 0) if now else 0; old_revenue = float(old.revenue or 0) if old else 0
        result.append({"category": (now or old).category, "article": (now or old).article, "code": (now or old).code, "name": (now or old).name, "revenue": revenue, "previous_revenue": old_revenue, "revenue_change": revenue - old_revenue, "units": float(now.units or 0) if now else 0, "checks": int(now.checks) if now else 0, "share": revenue / total * 100 if total else 0})
    return sorted(result, key=lambda row: row["revenue"], reverse=True)


@router.get("/{category_name}/details")
async def details(
    category_name: str,
    date_from: date,
    date_to: date,
    comparison_mode: str = Query(
        "previous_period", pattern="^(previous_period|previous_year)$"
    ),
    period_kind: str = "custom",
    stores: list[str] = Query(default=[]),
    manager: str | None = None,
    buyer_type: str | None = None,
    group_by: str | None = Query(None, pattern="^(day|week|month)$"),
    db: AsyncSession = Depends(get_db),
):
    start, end, old_start, old_end, _, clients, category_map, _, rows, _, _ = (
        await _scope(
            date_from,
            date_to,
            period_kind,
            stores,
            manager,
            buyer_type,
            category_name,
            db,
            comparison_mode=comparison_mode,
        )
    )
    if not rows:
        raise HTTPException(404, "Категория не найдена")
    products = await _product_details(
        db, start, end, old_start, old_end, stores, clients, category_map, category_name
    )
    category_map, join, category = _mapped_items(category_map)
    grouping = group_by or default_grouping(start, end)
    bucket = func.date_trunc(grouping, Sale.sale_date).label("period")
    points = (
        await db.execute(
            select(
                bucket,
                func.sum(SaleItem.quantity * SaleItem.actual_price).label("revenue"),
                func.sum(SaleItem.quantity).label("units"),
            )
            .select_from(SaleItem)
            .join(Sale)
            .outerjoin(category_map, join)
            .where(
                *_sale_conditions(start, end, stores, clients),
                category == category_name,
            )
            .group_by(bucket)
            .order_by(bucket)
        )
    ).all()
    return {
        "category": {key: value for key, value in rows[0].items() if key != "products"},
        "products": products,
        "group_by": grouping,
        "points": [
            {
                "period": row.period.date(),
                "revenue": float(row.revenue or 0),
                "units": float(row.units or 0),
            }
            for row in points
        ],
    }


def _sheet(book, title, headers, rows):
    sheet = book.create_sheet(title); sheet.append(headers)
    for row in rows: sheet.append(row)
    sheet.freeze_panes = "A2"
    for index, cells in enumerate(sheet.columns, 1): sheet.column_dimensions[get_column_letter(index)].width = min(45, max(12, max(len(str(cell.value or "")) for cell in cells) + 2))


@router.get("/export/xlsx")
async def export(date_from: date, date_to: date, period_kind: str = "custom", stores: list[str] = Query(default=[]), manager: str | None = None, buyer_type: str | None = None, category: str | None = None, db: AsyncSession = Depends(get_db)):
    start, end, old_start, old_end, _, clients, category_map, _, rows, checks, old_checks = await _scope(date_from, date_to, period_kind, stores, manager, buyer_type, category, db)
    metrics = report_summary(rows, checks, old_checks); products = await _product_details(db, start, end, old_start, old_end, stores, clients, category_map, None)
    book = Workbook(); book.remove(book.active)
    _sheet(book, "Итоги", ["Показатель", "Текущий", "Предыдущий", "Изменение", "Изменение, %"], [[name, value["current"], value["previous"], value["difference"], value["change_percent"]] for name, value in metrics.items()] + [["Период", str(start), str(end), str(old_start), str(old_end)]])
    headers = ["Категория", "Выручка", "Доля, %", "Продано", "Чеков", "Средний чек", "Прошлая выручка", "Изменение", "Изменение, %", "Изменение доли, п.п.", "Вклад, %"]
    category_rows = [[row["category"], row["current_revenue"], row["current_share"], row["current_units"], row["current_checks"], row["current_average_check"], row["previous_revenue"], row["revenue_change"], row["revenue_change_percent"], row["share_change_pp"], row["contribution_percent"]] for row in rows]
    _sheet(book, "Категории", headers, category_rows); _sheet(book, "Драйверы роста", headers, [row for row in category_rows if row[7] > 0]); _sheet(book, "Категории с падением", headers, [row for row in category_rows if row[7] < 0]); _sheet(book, "Товары", ["Категория", "Артикул", "Код", "Наименование", "Выручка", "Прошлая выручка", "Изменение", "Продано", "Чеков", "Доля"], [[row[key] for key in ("category", "article", "code", "name", "revenue", "previous_revenue", "revenue_change", "units", "checks", "share")] for row in products])
    output = BytesIO(); book.save(output)
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=category-analytics.xlsx"})


@router.get("/product-sales/list")
async def product_sales(
    date_from: date,
    date_to: date,
    period_kind: str = "custom",
    code: str | None = None,
    article: str | None = None,
    name: str | None = None,
    stores: list[str] = Query(default=[]),
    manager: str | None = None,
    buyer_type: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
):
    if date_from > date_to:
        raise HTTPException(422, "Дата начала не может быть позже даты окончания")
    start, end, _ = effective_period(
        date_from,
        date_to,
        period_kind,
        datetime.now(ZoneInfo(settings.autoload_timezone)).date(),
    )
    clients = await _clients(db, manager, buyer_type)
    if code and code.strip():
        product = func.lower(func.trim(SaleItem.code)) == code.strip().casefold()
    elif article and article.strip():
        product = and_(
            or_(SaleItem.code.is_(None), func.trim(SaleItem.code) == ""),
            func.lower(func.trim(SaleItem.article)) == article.strip().casefold(),
        )
    elif name and name.strip():
        product = and_(
            or_(SaleItem.code.is_(None), func.trim(SaleItem.code) == ""),
            or_(SaleItem.article.is_(None), func.trim(SaleItem.article) == ""),
            SaleItem.name == name,
        )
    else:
        raise HTTPException(422, "Не указан товар")
    matches = select(SaleItem.id).where(SaleItem.sale_id == Sale.id, product).exists()
    conditions = [*_sale_conditions(start, end, stores, clients), matches]
    total = int(
        await db.scalar(select(func.count()).select_from(Sale).where(*conditions)) or 0
    )
    rows = (
        await db.execute(
            select(
                Sale.id,
                Sale.sale_date,
                Sale.document_number,
                Sale.client,
                Sale.department,
                Sale.total_amount,
                Sale.base_amount,
                Sale.discount_percent,
                Sale.author,
                Sale.price_type,
            )
            .where(*conditions)
            .order_by(Sale.sale_date.desc(), Sale.id.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
    ).all()
    return {
        "items": [dict(row._mapping) for row in rows],
        "page": page,
        "pages": max(1, (total + page_size - 1) // page_size),
        "total": total,
    }
