from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from app.config import settings
from app.services.seasonality_storage import attach, metadata
from app.services.seasonality import STATUSES
from io import BytesIO
from itertools import islice
import time
import math
from app.services.catalog_sync import sync_image_url

from app.routers.product_analytics import HeavyAnalyticsRoute
from app.services.product_analytics_runtime import (
    diagnostics,
    available_memory,
    ensure_headroom,
    run_cpu,
)
from app.services.promotion_selected_export import SelectedExport, excel_selected
from app.services import vrcatalog

from fastapi import APIRouter, Depends, HTTPException, Response
from fastapi.routing import APIRoute
from openpyxl import Workbook
from pydantic import BaseModel, Field
from sqlalchemy import case, distinct, false, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models import Sale, SaleItem
from app.services.clients_vr import ClientsVrError
from app.services.product_analytics import catalog_property, product_key
from app.services.promotion_candidates import (
    NOT_FOUND,
    TABLE_ROW_LIMIT,
    UNCATEGORIZED,
    classify_candidate,
    percent_change,
    table_rows,
)
from app.services.sales_client_filters import get_sales_filter_clients
from app.services.vrcatalog import VrCatalogError, get_catalog_batch_info


class PromotionRoute(HeavyAnalyticsRoute):
    async def handle(self, scope, receive, send):
        if self.path.endswith("/history"):
            return await APIRoute.handle(self, scope, receive, send)
        return await super().handle(scope, receive, send)


router = APIRouter(
    route_class=PromotionRoute,
    prefix="/reports/promotion-candidates",
    tags=["Отчеты"],
)


class Criteria(BaseModel):
    decline_percent: float = -30
    min_previous_units: float = 5
    stale_days: int = 30
    low_sales_units: float = 2
    excess_stock_months: float = 6
    min_stock: float | None = 1
    recent_promotion_months: int = 3


class Request(BaseModel):
    date_from: date
    date_to: date
    departments: list[str] = Field(default_factory=list)
    manager: str | None = None
    buyer_type: str | None = None
    category: str | None = None
    brand: str | None = None
    manufacturer: str | None = None
    reason: str | None = None
    status: str | None = None
    use_seasonality: bool = True
    seasonality_status: str | None = None
    criteria: Criteria = Field(default_factory=Criteria)


def _conditions(data, clients):
    conditions = [Sale.sale_date >= data.date_from, Sale.sale_date <= data.date_to]
    if data.departments:
        conditions.append(Sale.department.in_(data.departments))
    if clients is not None:
        names = [x.strip().casefold() for x in clients if x.strip()]
        conditions.append(
            func.lower(func.trim(Sale.client)).in_(names) if names else false()
        )
    return conditions


async def _sales(data, db, clients):
    end = data.date_to
    last = end - timedelta(days=29)
    previous_start = end - timedelta(days=59)
    previous_end = end - timedelta(days=30)
    old_start = end - timedelta(days=89)
    old_end = end - timedelta(days=60)
    units = func.coalesce(func.sum(SaleItem.quantity), 0)
    revenue = func.coalesce(func.sum(SaleItem.quantity * SaleItem.actual_price), 0)

    def conditional(value, start, finish):
        return func.coalesce(
            func.sum(case((Sale.sale_date.between(start, finish), value), else_=0)), 0
        )

    result = await db.stream(
        select(
            SaleItem.code,
            SaleItem.article,
            func.max(SaleItem.name).label("name"),
            units.label("units"),
            revenue.label("revenue"),
            func.count(distinct(Sale.id)).label("checks"),
            conditional(SaleItem.quantity, last, end).label("units_30"),
            conditional(SaleItem.quantity, previous_start, previous_end).label(
                "units_previous_30"
            ),
            conditional(SaleItem.quantity, old_start, old_end).label("units_60_90"),
            conditional(SaleItem.quantity * SaleItem.actual_price, last, end).label(
                "revenue_30"
            ),
            conditional(
                SaleItem.quantity * SaleItem.actual_price, previous_start, previous_end
            ).label("revenue_previous_30"),
            func.min(Sale.sale_date).label("first_sale"),
            func.max(Sale.sale_date).label("last_sale"),
            func.count(distinct(Sale.department)).label("stores"),
        )
        .join(Sale)
        .where(*_conditions(data, clients))
        .group_by(SaleItem.code, SaleItem.article)
        .execution_options(yield_per=250)
    )
    rows = []
    try:
        async for x in result:
            if len(rows) % 250 == 0:
                ensure_headroom()
            rows.append(
                {
                    "key": product_key(x.article, x.code, x.name),
                    "code": x.code,
                    "article": x.article,
                    "name": x.name,
                    "units": float(x.units),
                    "revenue": float(x.revenue),
                    "checks": x.checks,
                    "average_price": (
                        float(x.revenue) / float(x.units) if x.units else 0
                    ),
                    "units_30": float(x.units_30),
                    "units_previous_30": float(x.units_previous_30),
                    "units_60_90": float(x.units_60_90),
                    "revenue_30": float(x.revenue_30),
                    "revenue_previous_30": float(x.revenue_previous_30),
                    "first_sale": x.first_sale,
                    "last_sale": x.last_sale,
                    "stores": x.stores,
                    "analysis_days": (
                        (data.date_to - x.first_sale).days + 1 if x.first_sale else 0
                    ),
                }
            )
    finally:
        await result.close()
    return rows


def _stock(info):
    values = (
        info.get("stocks") or info.get("warehouse_stocks") or info.get("balances") or []
    )
    items = (
        values.values()
        if isinstance(values, dict)
        else values if isinstance(values, list) else ()
    )
    result = 0.0
    found = False
    for item in items:
        raw = (
            item.get("quantity", item.get("stock_quantity", item.get("value")))
            if isinstance(item, dict)
            else item
        )
        try:
            result += float(raw)
            found = True
        except (TypeError, ValueError):
            pass
    return result if found else None


def _aviatorov_stock(info):
    """Read only the named warehouse; unknown is not an invented zero."""
    values = (
        info.get("stocks") or info.get("warehouse_stocks") or info.get("balances") or []
    )
    entries = (
        values.items()
        if isinstance(values, dict)
        else ((None, item) for item in values) if isinstance(values, list) else ()
    )
    total = 0.0
    found = False
    for key, item in entries:
        name = (
            (item.get("warehouse") or item.get("name") or key)
            if isinstance(item, dict)
            else key
        )
        if (
            not isinstance(name, str)
            or " ".join(name.split()).casefold() != "авиаторов зал+склад"
        ):
            continue
        raw = (
            item.get("quantity", item.get("stock_quantity", item.get("value")))
            if isinstance(item, dict)
            else item
        )
        try:
            quantity = float(raw)
            if not math.isfinite(quantity):
                continue
            total += quantity
            found = True
        except (TypeError, ValueError):
            pass
    return total if found else None


def _promotion_projection(info):
    image = sync_image_url(vrcatalog.catalog_item_image_url(info))
    return {
        "category": catalog_property(info, "category"),
        "brand": catalog_property(info, "brand"),
        "manufacturer": catalog_property(info, "manufacturer"),
        "stock": _stock(info),
        "aviatorov_stock": _aviatorov_stock(info),
        "image_url": image,
    }


async def _catalog(rows):
    # Keep legacy 5000-request boundaries and last-alias-wins merge order.
    products = (
        {"code": x["code"], "article": x["article"]}
        for x in rows
        if x["code"] or x["article"]
    )
    result = {}
    try:
        while batch := list(islice(products, 5000)):
            ensure_headroom()
            result.update(
                await get_catalog_batch_info(batch, project=_promotion_projection)
            )
    except VrCatalogError:
        return {}, False
    return result, True


def _stage(stage, started, **counts):
    diagnostics(
        "promotion_" + stage,
        elapsed_ms=(time.perf_counter() - started) * 1000,
        catalog_cache_charged_bytes=vrcatalog._catalog_info_bytes,
        catalog_cache_entries=len(vrcatalog.catalog_info_cache),
        available_bytes=available_memory(),
        **counts,
    )


async def _dataset(data, db):
    if data.seasonality_status and data.seasonality_status not in STATUSES:
        raise HTTPException(422, "Неизвестный статус сезонности")
    if data.date_from > data.date_to:
        raise HTTPException(422, "Дата начала не может быть позже даты окончания")
    try:
        clients = await get_sales_filter_clients(db, data.manager, data.buyer_type)
    except ClientsVrError as exc:
        raise HTTPException(502, str(exc)) from exc
    started = time.perf_counter()
    ensure_headroom()
    rows = await _sales(data, db, clients)
    _stage("sales", started, rows=len(rows))
    started = time.perf_counter()
    catalog, available = await _catalog(rows)
    _stage("catalog", started, rows=len(rows), catalog_keys=len(catalog))
    started = time.perf_counter()
    for index, row in enumerate(rows):
        if index % 250 == 0:
            ensure_headroom()
        info = catalog.get(row["key"])
        row.update(
            category=(
                info["category"]
                if info
                else NOT_FOUND if available else "CatalogVR недоступен"
            ),
            brand=info["brand"] if info else "Не заполнено",
            manufacturer=info["manufacturer"] if info else "Не заполнено",
            stock=info["stock"] if info else None,
            aviatorov_stock=info.get("aviatorov_stock") if info else None,
            image_url=info.get("image_url") if info else None,
        )
        if row["category"] == "Не заполнено":
            row["category"] = UNCATEGORIZED
    del catalog
    _stage("enriched", started, rows=len(rows))
    started = time.perf_counter()
    category_totals = {}
    for row in rows:
        target = category_totals.setdefault(row["category"], [0, 0])
        target[0] += row["units_30"]
        target[1] += row["units_previous_30"]
    criteria = data.criteria.model_dump()
    for index, row in enumerate(rows):
        if index % 250 == 0:
            ensure_headroom()
        current, previous = category_totals[row["category"]]
        row["category_change_percent"] = percent_change(current, previous)
        classify_candidate(row, criteria, data.date_to)
        row["stock_value"] = (
            row["stock"] * row["average_price"]
            if row["stock"] is not None and row["average_price"]
            else None
        )
    started = time.perf_counter()
    await attach(
        rows,
        db,
        criteria,
        datetime.now(ZoneInfo(settings.autoload_timezone)).date(),
        comparable_scope=not (data.departments or data.manager or data.buyer_type)
        and 0
        <= (
            datetime.now(ZoneInfo(settings.autoload_timezone)).date() - data.date_to
        ).days
        <= 1,
        analysis_end=data.date_to,
    )
    _stage("seasonality_read", started, rows=len(rows))
    kept = 0
    for x in rows:
        if (
            (not data.category or x["category"].casefold() == data.category.casefold())
            and (not data.brand or x["brand"].casefold() == data.brand.casefold())
            and (
                not data.manufacturer
                or x["manufacturer"].casefold() == data.manufacturer.casefold()
            )
            and (not data.reason or data.reason in x["reasons"])
            and (not data.status or x["status"] == data.status)
            and (
                not data.seasonality_status
                or x["seasonality"]["status"] == data.seasonality_status
            )
        ):
            rows[kept] = x
            kept += 1
    del rows[kept:]
    ensure_headroom()
    rows.sort(
        key=lambda x: ((x["seasonality"]["priority"],) if data.use_seasonality else ())
        + (len(x["reasons"]), x["stock_value"] or 0),
        reverse=True,
    )
    _stage("classified", started, rows=len(rows))
    return rows, available


@router.post("")
async def report(data: Request, db: AsyncSession = Depends(get_db)):
    rows, available = await _dataset(data, db)
    return {
        "items": table_rows(rows),
        "total": len(rows),
        "table_limit": TABLE_ROW_LIMIT,
        "summary": {
            "candidates": sum(x["status"] == "Кандидат" for x in rows),
            "stale": sum("Давно не продавался" in x["reasons"] for x in rows),
            "excess": sum("Избыточный запас" in x["reasons"] for x in rows),
            "decline": sum("Продажи падают" in x["reasons"] for x in rows),
        },
        "catalog_available": available,
        "seasonality": await metadata(
            db, datetime.now(ZoneInfo(settings.autoload_timezone)).date()
        ),
        "limitations": [
            "Историческое поле «Вид товара» не хранится в Sale/SaleItem; эффективность прошлых акций недоступна без изменения импорта."
        ],
    }


@router.post("/history")
async def history(_data: Request):
    return {
        "items": [],
        "available": False,
        "message": "История недоступна: «Вид товара» не сохраняется в строках продаж. Текущее свойство CatalogVR не позволяет восстановить исторические периоды акции.",
    }


@router.post("/export")
async def export(data: Request, db: AsyncSession = Depends(get_db)):
    rows, _ = await _dataset(data, db)
    return await run_cpu(_excel, rows)


def _excel(rows):
    started = time.perf_counter()
    book = Workbook(write_only=True)
    try:
        sheet = book.create_sheet("Кандидаты")
        sheet.append(
            [
                "Товар",
                "Код",
                "Артикул",
                "Категория",
                "Бренд",
                "Причины",
                "Статус",
                "Выручка",
                "Продано",
                "Продажи 30",
                "Предыдущие 30",
                "Изменение, %",
                "Категория, %",
                "Остаток",
                "Запас, мес.",
                "Дней без продажи",
                "Статус сезонности",
                "Сезонный индекс",
                "Сезонный приоритет",
                "Уверенность, баллы",
                "Ожидаемый период спроса",
                "Рекомендуемое действие",
                "Причина рекомендации",
                "Остатки Авиаторов Зал+Склад",
            ]
        )
        for index, x in enumerate(rows):
            if index % 250 == 0:
                ensure_headroom()
            season = x.get("seasonality") or {}
            sheet.append(
                [
                    x["name"],
                    x["code"],
                    x["article"],
                    x["category"],
                    x["brand"],
                    ", ".join(x["reasons"]),
                    x["status"],
                    x["revenue"],
                    x["units"],
                    x["units_30"],
                    x["units_previous_30"],
                    x["units_change_percent"],
                    x["category_change_percent"],
                    x["stock"],
                    x["stock_months"],
                    x["days_without_sales"],
                    season.get("status", "Недостаточно данных"),
                    season.get("index"),
                    season.get("priority", 0),
                    season.get("confidence", 0),
                    season.get("period", "—"),
                    season.get("action", "Сохранить оценку по действующим критериям"),
                    season.get(
                        "reason",
                        "Готовые актуальные показатели сезонности отсутствуют.",
                    ),
                    x.get("aviatorov_stock"),
                ]
            )
        history = book.create_sheet("История акций")
        history.append(
            ["История недоступна", "Поле «Вид товара» не сохраняется в SaleItem"]
        )
        output = BytesIO()
        book.save(output)
        ensure_headroom()
        _stage("excel", started, rows=len(rows), xlsx_bytes=output.tell())
        return Response(
            output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": "attachment; filename=promotion-candidates.xlsx"
            },
        )
    finally:
        # write-only XML temp files must be removed even if a memory check rejects.
        for sheet in book.worksheets:
            try:
                if not sheet.closed:
                    sheet.close()
                if sheet._writer is not None:
                    sheet._writer.cleanup()
            except (OSError, ValueError):
                pass
        book.close()


class ProductCatalogRequest(BaseModel):
    code: str | None = Field(default=None, max_length=500)
    article: str | None = Field(default=None, max_length=500)


def _product_card_projection(info):
    values = (
        info.get("stocks") or info.get("warehouse_stocks") or info.get("balances") or []
    )
    entries = (
        values.items()
        if isinstance(values, dict)
        else ((None, item) for item in values) if isinstance(values, list) else ()
    )
    stocks = []
    for key, item in entries:
        name = (
            (item.get("warehouse") or item.get("name") or key)
            if isinstance(item, dict)
            else key
        )
        if not isinstance(name, str) or not name.strip():
            continue
        raw = (
            item.get("quantity", item.get("stock_quantity", item.get("value")))
            if isinstance(item, dict)
            else item
        )
        try:
            quantity = float(raw)
            if not math.isfinite(quantity):
                quantity = None
        except (TypeError, ValueError):
            quantity = None
        stocks.append({"name": name.strip(), "quantity": quantity})
    manager = catalog_property(info, "Менеджер ОЗ")
    return {
        "stocks": stocks,
        "oz_manager": None if manager == "Не заполнено" else manager,
    }


@router.post("/product-catalog")
async def product_catalog(data: ProductCatalogRequest):
    if not (data.code and data.code.strip() or data.article and data.article.strip()):
        return {"stocks": [], "oz_manager": None}
    ensure_headroom()
    try:
        result = await get_catalog_batch_info(
            [data.model_dump()], project=_product_card_projection
        )
    except VrCatalogError:
        raise HTTPException(502, "Не удалось загрузить данные VRCatalog") from None
    info = result.get(product_key(data.article, data.code, ""))
    if info is None:
        raise HTTPException(502, "Данные товара не получены от VRCatalog")
    return info


@router.post("/export-selected")
async def export_selected(data: SelectedExport):
    return await run_cpu(excel_selected, data)
